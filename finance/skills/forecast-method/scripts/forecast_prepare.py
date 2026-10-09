# /// script
# dependencies = ["pyyaml", "openpyxl"]
# ///
"""Open or refresh a forecast vintage before a session, reading the workbooks only.

In order:

1. Which two forecasts. With --vintage naming an existing vintage, the workbooks its
   VINTAGE.yaml records. Otherwise --new (or the newest issued revision in the settings'
   revisions.folder matching revisions.pattern, ordered by the number revisions.number captures,
   then by modification time) against --prior (or the newest vintage's new workbook, or the
   revision issued before). With nothing newer than the newest vintage it prints NOTHING:.
   A rolling vintage (--kind rolling) needs --vintage and --last-actual-month; its new workbook
   is built later by forecast-revise.
2. The vintage folder. vintages/<V>/VINTAGE.yaml is written once and never rewritten (a person
   may edit it); a different --new for an existing vintage is refused.
3. The extracts. work/source/prior.json and new.json (and budget.json when the vintage names a
   budget_workbook), each footed and tied. An extract of a workbook that changed is moved to
   work/source/superseded-<stamp>/, never overwritten.
4. The evidence index, work/source/evidence.json: one E-nnn per non-empty row of the new
   workbook's change_log_sheets and per file in evidence_dirs. An id keeps its meaning; new
   items get the next number.
5. work/source/pulled.md: when, which files, their SHA-256, and whether they foot and tie.

The first line printed is FRESH:, STALE: (an extract does not foot or tie) or NOTHING:. Exit 0
whenever it ran, 2 on a bad argument or an unreadable workbook. --dry-run writes nothing.

Example:
    python3 forecast_prepare.py ~/Forecast
    python3 forecast_prepare.py ~/Forecast --vintage 2026-10 --kind rolling --prior "delivery/v1.xlsx" --last-actual-month 2026-09
"""

import argparse
import json
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

from _common import (VINTAGE_FILE, ForecastError, current_vintage, extract, forecast_folder, ident, load_settings,
                     load_vintage, now, read_grid, resolve_path, revisions, vintage_dir, write_json)


def slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "-", str(text)).strip("-").lower()[:64].strip("-") or "vintage"


def vintage_id_for(settings, workbook):
    """A vintage id (and extract label) from a workbook name, via revisions.vintage_from when set."""
    pattern = (settings.get("revisions") or {}).get("vintage_from")
    match = re.search(pattern, workbook.stem, re.I) if pattern else None
    return slug(match.group(0) if match else workbook.stem)


def choose(root, settings, vintage, new, prior, kind):
    """Which vintage, which two workbooks, and whether there is anything to do."""
    if vintage and (vintage_dir(root, vintage) / VINTAGE_FILE).is_file():
        meta = load_vintage(root, vintage)
        if new and resolve_path(root, new) != resolve_path(root, meta.get("new_workbook", "")):
            raise ForecastError(f"vintage {vintage} already records new workbook {meta.get('new_workbook')}; "
                                "start a new vintage for a different workbook")
        return {"vintage": vintage, "existing": True, "meta": meta,
                "new": resolve_path(root, meta["new_workbook"]), "prior": resolve_path(root, meta["prior_workbook"])}
    latest = current_vintage(root)
    latest_meta = load_vintage(root, latest) if latest else {}
    carried = {"prior_vintage": latest or "", "prior_last_actual_month": latest_meta.get("last_actual_month", "")}
    if kind == "rolling":
        if not vintage:
            raise ForecastError("a rolling vintage needs --vintage (the month it is built, such as 2026-10)")
        prior_path = resolve_path(root, prior) if prior else (
            resolve_path(root, latest_meta["new_workbook"]) if latest_meta else None)
        if prior_path is None or not prior_path.is_file():
            raise ForecastError("a rolling vintage rolls forward from a prior workbook: name --prior")
        pattern = str((settings.get("revise") or {}).get("name") or "{vintage} forecast DRAFT.xlsx")
        new_path = resolve_path(root, new) if new else \
            vintage_dir(root, vintage) / "revision" / pattern.format(vintage=vintage, prior=prior_path.stem)
        return dict(carried, vintage=vintage, existing=False, meta=None, new=new_path, prior=prior_path)
    issued = revisions(settings)
    if new:
        new_path = resolve_path(root, new)
    elif issued:
        new_path = issued[-1]
        if latest_meta and resolve_path(root, latest_meta.get("new_workbook", "")) == new_path:
            return {"nothing": f"the newest revision {new_path.name} is already vintage {latest}"}
    else:
        raise ForecastError("name --new, or set revisions.folder and revisions.pattern in the settings")
    if prior:
        prior_path = resolve_path(root, prior)
    elif latest_meta:
        prior_path = resolve_path(root, latest_meta["new_workbook"])
    else:
        earlier = [p for p in issued if p != new_path]
        if not earlier:
            raise ForecastError("no prior forecast: name --prior")
        before = [p for p in earlier if issued.index(p) < issued.index(new_path)] if new_path in issued else earlier
        prior_path = (before or earlier)[-1]
    for path in (new_path, prior_path):
        if not path.is_file():
            raise ForecastError(f"workbook does not exist: {path}")
    return dict(carried, vintage=vintage or vintage_id_for(settings, new_path), existing=False, meta=None,
                new=new_path, prior=prior_path)


def change_log_items(workbook, sheets):
    """One evidence item per non-empty row of each change-log sheet the workbook has."""
    from openpyxl import load_workbook
    book = load_workbook(str(workbook), read_only=True)
    present = set(book.sheetnames)
    book.close()
    items = []
    for sheet in (s for s in sheets if s in present):
        for number, row in enumerate(read_grid(workbook, sheet), start=1):
            text = " | ".join(c for c in (ident(v, "") for v in row if v not in (None, "")) if c)
            if len(text) >= 3:
                items.append({"kind": "change-log", "source": f"{Path(workbook).name}!{sheet} row {number}",
                              "text": text[:600]})
    return items


def file_items(root, settings):
    """One evidence item per file in evidence_dirs, modified on or after evidence_since."""
    since = str(settings.get("evidence_since") or "")
    items = []
    for entry in settings.get("evidence_dirs") or []:
        folder = resolve_path(root, entry)
        for path in sorted(folder.rglob("*")) if folder.is_dir() else []:
            if not path.is_file() or path.name.startswith(("~$", ".")):
                continue
            modified = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
            if not since or modified >= since:
                items.append({"kind": "file", "source": str(path), "date": modified, "size": path.stat().st_size})
    return items


def index_evidence(existing, found):
    """Merge found items into the index: an id keeps its source; new items get the next number."""
    by_source = {(i["kind"], i["source"]): i for i in existing}
    following = max((int(i["id"][2:]) for i in existing if re.match(r"^E-\d+$", str(i.get("id", "")))), default=0)
    out = list(existing)
    for item in found:
        key = (item["kind"], item["source"])
        if key in by_source:
            by_source[key].update({k: v for k, v in item.items() if k != "id"})
            continue
        following += 1
        by_source[key] = dict(item, id=f"E-{following:03d}")
        out.append(by_source[key])
    return out


def keep_superseded(source, name, workbook_sha, stamp):
    """Move an extract of a different workbook aside; return where, or None if it was current."""
    path = source / name
    if not path.is_file():
        return None
    try:
        old = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        old = {}
    if old.get("sha256") == workbook_sha:
        return None
    target = source / f"superseded-{stamp}"
    target.mkdir(parents=True, exist_ok=True)
    shutil.move(str(path), str(target / name))
    return str(target / name)


def prepare(folder, vintage="", new="", prior="", kind="revision", last_actual="", years=None, dry_run=False):
    import yaml
    root = forecast_folder(folder)
    settings = load_settings(root)
    pick = choose(root, settings, vintage, new, prior, kind)
    if "nothing" in pick:
        return {"line": f"NOTHING: {pick['nothing']}", "wrote": []}
    vid = pick["vintage"]
    vdir = vintage_dir(root, vid)
    meta = pick["meta"] or {
        "vintage": vid, "kind": kind, "created": now(), "prior_vintage": pick["prior_vintage"],
        "prior_workbook": str(pick["prior"]), "new_workbook": str(pick["new"]), "layout": "",
        "years": years or list(settings.get("years") or []),
        "last_actual_month": last_actual or str(settings.get("last_actual_month") or ""),
        "prior_last_actual_month": pick["prior_last_actual_month"] or (
            str(settings.get("last_actual_month") or "") if kind == "revision" else ""),
        "budget_workbook": str(settings.get("budget_workbook") or ""), "requested_by": "", "request": ""}
    if meta.get("kind") == "rolling" and not meta.get("last_actual_month"):
        raise ForecastError("a rolling vintage needs --last-actual-month")
    layout_name = str(meta.get("layout") or "")
    year_list = [int(y) for y in meta.get("years") or []] or None
    built = Path(pick["new"]).is_file()   # a rolling vintage's new workbook may not exist yet
    new_x = extract(pick["new"], root, layout_name, year_list, vintage_id_for(settings, pick["new"])) if built else None
    prior_x = extract(pick["prior"], root, layout_name, year_list, vintage_id_for(settings, pick["prior"]))
    budget_x = None
    if meta.get("budget_workbook"):
        budget_x = extract(resolve_path(root, meta["budget_workbook"]), root,
                           str(settings.get("budget_layout") or layout_name), year_list, label="budget")

    source = vdir / "work" / "source"
    found = (change_log_items(pick["new"], list(settings.get("change_log_sheets") or [])) if built else []) \
        + file_items(root, settings)
    index_path = source / "evidence.json"
    existing = json.loads(index_path.read_text(encoding="utf-8")).get("items", []) if index_path.is_file() else []
    evidence = index_evidence(existing, found)

    bad = [f"{x['label']}: {c['name']}" for x in (prior_x, new_x, budget_x) if x for c in x["checks"] if not c["ok"]]
    target = new_x["label"] if new_x else f"{Path(pick['new']).name} (not built yet: forecast-revise builds it)"
    line = (f"{'STALE' if bad else 'FRESH'}: vintage {vid}, {prior_x['label']} to {target}, "
            f"{len(evidence)} evidence item(s)" + (f"; does not foot or tie: {'; '.join(bad[:3])}" if bad else ""))
    result = {"line": line, "vintage": vid, "dir": str(vdir), "prior": str(pick["prior"]), "new": str(pick["new"]),
              "existing": pick["existing"], "evidence": len(evidence), "problems": bad, "wrote": []}
    if dry_run:
        return dict(result, line="DRY RUN " + line)

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    source.mkdir(parents=True, exist_ok=True)
    if not (vdir / VINTAGE_FILE).is_file():
        (vdir / VINTAGE_FILE).write_text(yaml.safe_dump(meta, sort_keys=False, allow_unicode=True), encoding="utf-8")
        result["wrote"].append(str(vdir / VINTAGE_FILE))
    for name, x in (("prior.json", prior_x), ("new.json", new_x), ("budget.json", budget_x)):
        if x is None:
            continue
        moved = keep_superseded(source, name, x["sha256"], stamp)
        if moved:
            result.setdefault("superseded", []).append(moved)
        write_json(source / name, x)
        result["wrote"].append(str(source / name))
    write_json(index_path, {"items": evidence})
    result["wrote"].append(str(index_path))
    pulled = [f"# Pulled {datetime.now().strftime('%Y-%m-%d %H:%M')}", "", line, ""]
    pulled += [f"- {name}: `{x['workbook']}` sha256 {x['sha256']} ({'foots and ties' if x['ok'] else 'FAILS'})"
               for name, x in (("prior", prior_x), ("new", new_x), ("budget", budget_x)) if x]
    (source / "pulled.md").write_text("\n".join(pulled) + "\n", encoding="utf-8")
    result["wrote"].append(str(source / "pulled.md"))
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description="Open or refresh a forecast vintage: extracts and evidence index.")
    parser.add_argument("folder", help="the Forecast folder")
    parser.add_argument("--vintage", default="", help="an existing vintage to refresh, or the id for a new one")
    parser.add_argument("--new", default="", help="the new forecast workbook; blank: the newest issued revision")
    parser.add_argument("--prior", default="", help="the prior forecast workbook; blank: the newest vintage's new one")
    parser.add_argument("--kind", choices=["revision", "rolling"], default="revision")
    parser.add_argument("--last-actual-month", default="", help="YYYY-MM, the last closed month in the new forecast")
    parser.add_argument("--years", default="", help="fiscal years to bridge, comma-separated; blank: the settings' years")
    parser.add_argument("--dry-run", action="store_true", help="report what it would do; write nothing")
    parser.add_argument("--format", choices=["text", "json"], default="text")
    args = parser.parse_args(argv)
    try:
        years = [int(y) for y in args.years.split(",") if y.strip()] or None
        result = prepare(args.folder, args.vintage.strip(), args.new.strip(), args.prior.strip(), args.kind,
                         args.last_actual_month.strip(), years, args.dry_run)
    except (ForecastError, ValueError) as exc:
        print(f"forecast-prepare: {exc}", file=sys.stderr)
        return 2
    if args.format == "json":
        print(json.dumps(result, indent=1))
    else:
        print(result["line"])
        for path in result["wrote"]:
            print(f"  wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
