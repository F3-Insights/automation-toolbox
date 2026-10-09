"""The board-package scripts against an invented company, Acme Components.

A close folder with a results file and two trial balance pulls, a Reporting folder, the rules,
a Context and owner settings; then the whole loop: pack, render the deck, start the figure
ledger, source it, tie it out, record the reviews, ask the owner, and the check says done. A
figure one unit off must stop it.
"""

import csv
import io
import importlib.util
import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parents[1]
# Loaded under its own name, so this skill's tests can run beside another skill's _common.
_spec = importlib.util.spec_from_file_location("board_package_common", SCRIPTS / "_common.py")
bc = sys.modules["board_package_common"] = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(bc)

PERIOD = "2026-09"
NAME = "acme-board"
DECK = f"Acme Board Package {PERIOD} DRAFT.pptx"
SCRIPT_MD = f"Acme Board Package {PERIOD} commentary script DRAFT.md"
EMAIL_MD = f"Acme Board Package {PERIOD} cover email DRAFT.md"

RESULTS_HTML = """<!doctype html><html><head><title>Acme Components September 2026 Results v2</title>
<style>td{color:red}</style></head><body>
<h1>Acme Components, September 2026 Results (v2)</h1>
<p>September EBITDA of $(50)k was $(20)k below forecast. Prepared 2026-10-06.</p>
<table><tr><th>$k</th><th>ACT</th><th>FCST</th><th>&Delta; FCST</th></tr>
<tr><td>Revenue</td><td>1,000</td><td>1,100</td><td>(100)</td></tr>
<tr><td>Gross profit</td><td>400</td><td>460</td><td>(60)</td></tr>
<tr><td>EBITDA</td><td>(50)</td><td>(30)</td><td>(20)</td></tr></table>
<table><tr><th>September revenue ($)</th><th>ACT</th></tr>
<tr><td>Widgets</td><td>700,250</td></tr><tr><td>Gadgets</td><td>299,750</td></tr>
<tr><td>Total</td><td>1,000,000</td></tr></table>
<p>Composition: widgets $700.3k and gadgets $299.8k.</p>
</body></html>
"""

RULES = """# Board rules (Acme Components)

## Package inputs

- Results file: `Acme {Month} {yyyy} Results*.html`
- Results unit: k
- Deck file: `Acme Board Package {yyyy-mm} DRAFT.pptx`
- Script file: `Acme Board Package {yyyy-mm} commentary script DRAFT.md`
- Cover email: `Acme Board Package {yyyy-mm} cover email DRAFT.md`
- Lender pack: none
- Previous packages: `*Monthly Financial Report ({yyyy-mm})*.pdf`
- Supporting files: `*ACT vs FCST*`
- Close gate: attested
- Closed by hand: 2026-09
- Banned phrases: game changer
"""

DECK_SPEC = {"slides": [
    {"layout": "title", "title": "Acme Components", "subtitle": "September 2026 board package"},
    {"title": "September revenue of $1.0M, $(100)k below forecast",
     "bullets": ["Gross margin 40.0%", ["Cash at month end $500k, up $120k in the month"]],
     "table": {"columns": ["$k", "ACT", "FCST"], "rows": [["Revenue", "1,000", "1,100"], ["EBITDA", "(50)", "(30)"]]},
     "notes": "Revenue was $1.0M against $1.1M forecast."},
]}

SOURCES = {
    "$1.0M": "results:Revenue | ACT", "$(100)k": "results:Revenue | Δ FCST", "40.0%": "calc:F_GP/F_REV*100",
    "$500k": "tb:10400", "$120k": "tb-month:10400", "1,000": "-tb-month:40000-49999",
    "1,100": "results:Revenue | FCST", "(50)": "results:EBITDA | ACT", "(30)": "results:EBITDA | FCST",
    "$1.1M": "results:R2", "$1,000.0k": "text:",
}


def tb(path: Path, period: str, rows) -> None:
    path.mkdir(parents=True, exist_ok=True)
    (path / f"trial-balance-{period}.json").write_text(json.dumps(
        {"as_of": f"{period}-30", "fiscal_year_start": "2026-01-01", "rows": rows}), encoding="utf-8")
    (path / "pulled.md").write_text(f"# Pulled data, {period} (FINAL)\n", encoding="utf-8")


def acct(account, ending, ytd):
    return {"Account": account, "Name": account, "Ending balance": ending, "YTD activity": ytd}


@pytest.fixture()
def acme(tmp_path, monkeypatch):
    contexts, rules, package = tmp_path / "contexts", tmp_path / "rules", tmp_path / "Reporting"
    close = tmp_path / "Month-End"
    for folder in (contexts, rules, package):
        folder.mkdir()
    (rules / bc.RULES_MD).write_text(RULES, encoding="utf-8")
    month = close / "2026" / PERIOD
    month.mkdir(parents=True)
    (month / "Acme September 2026 Results v1.html").write_text(
        RESULTS_HTML.replace("1,000</td><td>1,100", "999</td><td>1,100"), encoding="utf-8")
    (month / "Acme September 2026 Results v2.html").write_text(RESULTS_HTML, encoding="utf-8")
    (month / "Acme Sep ACT vs FCST netted.md").write_text("netted\n", encoding="utf-8")
    tb(month / "work" / "source", PERIOD, [acct("10400", 500000.0, 120000.0), acct("40000", -9000000.0, -9000000.0),
                                          acct("50000", 5400000.0, 5400000.0)])
    tb(close / "2026" / "2026-08" / "work" / "source", "2026-08",
       [acct("10400", 380000.0, 0.0), acct("40000", -8000000.0, -8000000.0), acct("50000", 4800000.0, 4800000.0)])
    (contexts / f"{NAME}.yaml").write_text(json.dumps({
        "name": NAME, "description": "Acme Components board package",
        "sources": [{"name": "rules", "kind": "folder", "path": str(rules)},
                    {"name": "close", "kind": "folder", "path": str(close)},
                    {"name": "package", "kind": "folder", "path": str(package)}]}), encoding="utf-8")
    settings = tmp_path / "settings.toml"
    settings.write_text(f'[board-package-workstream]\ncontexts_dir = "{contexts}"\n', encoding="utf-8")
    monkeypatch.setenv("F3I_TOOLBOX_SETTINGS", str(settings))
    return {"rules": rules, "package": package, "month": month, "period_dir": package / PERIOD}


def run(script, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / f"{script}.py"), *[str(a) for a in args]],
                          capture_output=True, text=True, env=os.environ.copy())


def check(*extra):
    res = run("board_package_check", NAME, "--period", PERIOD, "--format", "json", *extra)
    assert res.returncode == 0, res.stderr
    return json.loads(res.stdout)


def write_ledger(path, rows):
    path.unlink(missing_ok=True)
    path.write_text(bc.to_csv(rows, bc.PACKAGE_COLUMNS), encoding="utf-8")


def build_package(acme, tweak=None) -> Path:
    """Pack, render the deck, write the script and cover email, start and source the ledger."""
    res = run("board_package_pack", NAME, "--period", PERIOD)
    assert res.returncode == 0, res.stderr
    work = acme["period_dir"] / "work"
    data = json.loads(json.dumps(DECK_SPEC))
    if tweak:
        tweak(data)
    (work / "deck-spec.json").write_text(json.dumps(data), encoding="utf-8")
    deck = acme["period_dir"] / DECK
    res = run("board_package_render", work / "deck-spec.json", "--out", deck)
    assert res.returncode == 0, res.stderr
    (acme["period_dir"] / SCRIPT_MD).write_text(
        "# Commentary\n\nSeptember revenue was $1,000.0k, and EBITDA closed at $(50)k.\n", encoding="utf-8")
    (acme["period_dir"] / EMAIL_MD).write_text(
        "Dear board,\n\nThe September package is attached. Revenue was $1.0M.\n", encoding="utf-8")
    ledger = work / bc.PACKAGE_FIGURES_CSV
    res = run("board_package_figures", deck, acme["period_dir"] / SCRIPT_MD, acme["period_dir"] / EMAIL_MD,
              "--ledger", "--out", ledger)
    assert res.returncode == 0, res.stderr
    rows = list(csv.DictReader(io.StringIO(ledger.read_text(encoding="utf-8"))))
    ids = {}
    for r in rows:
        if r["printed"] == "(50)" and r["file"].endswith(".md"):
            r["source"], r["unit"] = "results:EBITDA | ACT", "k"
        else:
            r["source"] = SOURCES.get(r["printed"], "results:EBITDA | ACT")
        if r["printed"] in ("1,000", "1,100", "(50)", "(51)", "(30)") and r["file"].endswith(".pptx"):
            r["unit"] = "k"
        ids.setdefault(r["printed"], r["id"])
    gp = next(r for r in bc.read_csv(work / bc.RESULTS_FIGURES_CSV) if r["metric"] == "Gross profit | ACT")
    helper = f"F{len(rows) + 1}"
    rows.append({"id": helper, "file": "", "location": "", "printed": "400", "unit": "k",
                 "metric": "Gross profit (helper)", "source": f"results:{gp['id']}", "note": ""})
    for r in rows:
        if r["source"].startswith("calc:"):
            r["source"] = f"calc:{helper}/{ids['1,000']}*100"
    write_ledger(ledger, rows)
    return ledger


def finish(acme, tied=True):
    """Tie out, record both reviews, ask the owner."""
    res = run("board_package_tieout", NAME, "--period", PERIOD)
    assert res.returncode == (0 if tied else 1), res.stdout + res.stderr
    work = acme["period_dir"] / "work"
    for kind, note in (("finance", bc.FINANCE_REVIEW_MD), ("redteam", bc.REDTEAM_MD)):
        (work / note).write_text("PASS\n", encoding="utf-8")
        res = run("board_package_record", NAME, "review", "--period", PERIOD, "--kind", kind, "--state", "passed",
                  "--file", f"work/{note}")
        assert res.returncode == 0, res.stderr
    (work / bc.CONFIRMATIONS_MD).write_text("# Confirmations\n\n## CR-0123456789\n\n- Asked of: owner\n",
                                            encoding="utf-8")


# --------------------------------------------------------------------------- printed numbers

@pytest.mark.parametrize("text,value,decimals,suffix,percent", [
    ("$(74)k", -74, 0, "k", False), ("(507)", -507, 0, "", False), ("+137", 137, 0, "", False),
    ("1,014,828", 1014828, 0, "", False), ("$104.3k", 104.3, 1, "k", False), ("40.2%", 40.2, 1, "", True),
    ("\u221274.5", -74.5, 1, "", False), ("$1.2M", 1.2, 1, "m", False), ("($74k)", -74, 0, "k", False),
])
def test_printed_numbers_parse_as_a_reader_reads_them(text, value, decimals, suffix, percent):
    p = bc.parse_printed(text)
    assert p is not None and p.value == value and p.decimals == decimals
    assert p.suffix == suffix and p.percent is percent


def test_running_text_counts_only_marked_numbers():
    found = [p.text for p in bc.tokens("Entity 11 had 3 sites on 2026-09-14 (v1); EBITDA $(74)k (+$104k), margin 40.2%")
             if bc.figure_in_text(p)]
    assert found == ["$(74)k", "+$104k", "40.2%"]


def test_rounding_to_the_printed_precision():
    assert bc.round_to_print(950808, bc.parse_printed("951"), "k") == 951000
    q = bc.parse_printed("$1.0M")
    assert bc.round_to_print(1049999, q, None) == 1000000
    assert bc.round_to_print(1050000, q, None) == 1100000


# --------------------------------------------------------------------------- figures

def test_results_figures_name_row_and_column_and_carry_the_table_unit(acme):
    out = acme["month"].parent / "r.csv"
    res = run("board_package_figures", acme["month"] / "Acme September 2026 Results v2.html", "--results",
              "--period", PERIOD, "--unit", "1", "--out", out)
    assert res.returncode == 0, res.stderr
    by_metric = {r["metric"]: r for r in bc.read_csv(out)}
    assert by_metric["Revenue | ACT"]["value"] == "1000000.0"
    assert by_metric["EBITDA | ACT"]["value"] == "-50000.0"
    assert by_metric["Total | ACT"]["value"] == "1000000.0" and by_metric["Total | ACT"]["unit"] == "1"


def test_figures_out_refuses_to_overwrite(acme, tmp_path):
    html = acme["month"] / "Acme September 2026 Results v2.html"
    assert run("board_package_figures", html, "--out", tmp_path / "f.csv").returncode == 0
    res = run("board_package_figures", html, "--out", tmp_path / "f.csv")
    assert res.returncode == 1 and "never overwrites" in res.stderr


def test_figures_reads_a_workbook_by_cell(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "TTM"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "Month", 2026, "Revenue", 1234.5
    ws["B2"].number_format = "#,##0.0"
    wb.save(tmp_path / "lender.xlsx")
    found = bc.extract(tmp_path / "lender.xlsx")
    assert [(f["location"], f["printed"]) for f in found] == [("TTM!B2", "1234.5")]


# --------------------------------------------------------------------------- pack

def test_pack_reads_the_newest_results_version_and_the_close(acme):
    res = run("board_package_pack", NAME, "--period", PERIOD)
    assert res.returncode == 0 and res.stdout.startswith("READY: acme-board 2026-09"), res.stdout + res.stderr
    pack = json.loads((acme["period_dir"] / "work" / bc.SOURCES_JSON).read_text(encoding="utf-8"))
    assert pack["results"]["file"]["name"] == "Acme September 2026 Results v2.html"
    assert pack["close"]["done"] is True and pack["pulls"][0]["label"] == "FINAL"
    assert [s["name"] for s in pack["supporting"]] == ["Acme Sep ACT vs FCST netted.md"]


def test_pack_on_a_dry_run_writes_only_under_the_run_folder(acme, tmp_path):
    run_dir = tmp_path / "run"
    res = run("board_package_pack", NAME, "--period", PERIOD, "--run-dir", run_dir, "--dry-run-if", "true")
    assert res.returncode == 0, res.stderr
    assert (run_dir / "board-package" / PERIOD / "work" / bc.SOURCES_JSON).is_file()
    assert not acme["period_dir"].exists()
    assert run("board_package_pack", NAME, "--period", PERIOD, "--dry-run-if", "true").returncode == 2


def test_pack_says_not_ready_when_the_close_is_open(acme):
    (acme["rules"] / bc.RULES_MD).write_text(RULES.replace("- Closed by hand: 2026-09\n", ""), encoding="utf-8")
    res = run("board_package_pack", NAME, "--period", PERIOD)
    assert res.stdout.startswith("NOT READY:") and "not in the rules' Closed by hand" in res.stdout


def test_pack_with_the_month_end_gate_and_no_finished_close_is_not_ready(acme, tmp_path, monkeypatch):
    monkeypatch.setenv("HOME", str(tmp_path))      # no month-end check is installed here
    (acme["rules"] / bc.RULES_MD).write_text(
        RULES.replace("- Close gate: attested\n- Closed by hand: 2026-09\n", ""), encoding="utf-8")
    res = run("board_package_pack", NAME, "--period", PERIOD)
    assert res.returncode == 0, res.stderr
    assert res.stdout.startswith("NOT READY:") and "close is not done" in res.stdout


def test_a_blank_period_is_the_month_before(acme):
    res = run("board_package_pack", NAME, "--period=", "--as-of", "2026-10-04")
    assert res.returncode == 0 and "2026-09" in res.stdout


def test_a_missing_contexts_setting_names_the_setting(acme, tmp_path, monkeypatch):
    monkeypatch.setenv("F3I_TOOLBOX_SETTINGS", str(tmp_path / "none.toml"))
    res = run("board_package_pack", NAME, "--period", PERIOD)
    assert res.returncode == 2 and "contexts_dir" in res.stderr


# --------------------------------------------------------------------------- render

def test_render_never_overwrites_and_checks_the_spec(tmp_path):
    spec, out = tmp_path / "spec.json", tmp_path / "deck.pptx"
    spec.write_text(json.dumps(DECK_SPEC), encoding="utf-8")
    assert run("board_package_render", spec, "--out", out).returncode == 0
    assert run("board_package_render", spec, "--out", out).returncode == 1
    res = run("board_package_render", spec, "--out", out, "--supersede", tmp_path / "work")
    assert res.returncode == 0 and out.is_file()
    assert len(list((tmp_path / "work").glob("superseded-*/deck.pptx"))) == 1
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"slides": [{"title": "T", "table": {"columns": ["a", "b"], "rows": [["1"]]}}]}),
                   encoding="utf-8")
    assert run("board_package_render", bad, "--out", tmp_path / "x.pptx").returncode == 2
    found = [f["printed"] for f in bc.extract(out)]
    assert "$1.0M" in found and "1,100" in found and "40.0%" in found


# --------------------------------------------------------------------------- tie-out, record and check

def test_the_whole_loop_is_done(acme):
    build_package(acme)
    finish(acme)
    result = check()
    assert {k: t["state"] for k, t in result["tests"].items()} == {k: "pass" for k in result["tests"]}, \
        json.dumps(result["tests"], indent=1)
    assert len(result["tests"]) == 9 and result["done"] is True
    rows = bc.read_csv(acme["period_dir"] / f"tieout-ledger-{PERIOD}.csv")
    assert {r["file"] for r in rows} >= {"books"}
    res = run("board_package_check", NAME, "--period", PERIOD, "--precheck")
    assert res.stdout.startswith("NOTHING: acme-board 2026-09: the package is drafted, tied and reviewed")
    shown = json.loads(run("board_package_record", NAME, "show", "--period", PERIOD, "--format", "json").stdout)
    assert {r["id"] for r in shown} == {"review:finance", "review:redteam"}


def test_a_figure_one_unit_off_stops_the_tieout_until_the_owner_approves(acme):
    def tweak(data):
        data["slides"][1]["table"]["rows"][1][1] = "(51)"
    ledger = build_package(acme, tweak)
    rows = bc.read_csv(ledger)
    for r in rows:
        if r["printed"] == "(51)":
            r["source"], r["unit"] = "results:EBITDA | ACT", "k"
    write_ledger(ledger, rows)
    res = run("board_package_tieout", NAME, "--period", PERIOD)
    assert res.returncode == 1 and "DIFFERS" in res.stdout
    # The disagreement names both sides with the place each figure came from.
    assert f"{DECK} -51,000 (slide 2 table), books -50,000 (results:EBITDA | ACT)" in res.stdout
    assert check()["tests"]["tieout"]["state"] == "fail"
    bad = next(r["id"] for r in rows if r["printed"] == "(51)")
    res = run("board_package_record", NAME, "exception", "--period", PERIOD, "--figure", bad,
              "--evidence", "CR-0123456789", "--note", "rounding agreed")
    assert res.returncode == 1 and "only the owner's answer" in res.stderr
    finish(acme, tied=False)
    res = run("board_package_record", NAME, "exception", "--period", PERIOD, "--figure", bad,
              "--evidence", "CR-0123456789", "--note", "the owner keeps (51): a late entry")
    assert res.returncode == 0, res.stderr
    assert run("board_package_tieout", NAME, "--period", PERIOD).returncode == 0
    assert check()["tests"]["tieout"]["state"] == "pass"


def test_a_figure_missing_from_the_ledger_fails_figures(acme):
    ledger = build_package(acme)
    write_ledger(ledger, [r for r in bc.read_csv(ledger) if r["printed"] != "$500k"])
    test = check()["tests"]["figures"]
    assert test["state"] == "fail" and "500k printed 1 time(s) more than the ledger lists" in test["why"]


def test_a_review_older_than_the_package_does_not_count(acme):
    build_package(acme)
    finish(acme)
    later = time.time() + 120
    os.utime(acme["period_dir"] / DECK, (later, later))
    tests = check()["tests"]
    assert tests["review"]["state"] == "fail" and "review again" in tests["review"]["why"]
    assert tests["redteam"]["state"] == "fail"


def test_a_person_stated_figure_must_be_named_in_the_review_note(acme):
    ledger = build_package(acme)
    rows = bc.read_csv(ledger)
    target = next(r for r in rows if r["printed"] == "$500k")
    target["source"] = "person:the controller, bank statement"
    write_ledger(ledger, rows)
    finish(acme)
    test = check()["tests"]["tieout"]
    assert test["state"] == "fail" and target["id"] in test["why"]
    (acme["period_dir"] / f"{PERIOD} package review notes.md").write_text(
        f"- {target['id']}: cash stated by the controller from the bank statement.\n", encoding="utf-8")
    assert check()["tests"]["tieout"]["state"] == "pass"


def test_prose_finds_banned_characters_and_phrases(acme):
    build_package(acme)
    script = acme["period_dir"] / SCRIPT_MD
    script.write_text(script.read_text(encoding="utf-8") + "\nA game changer \u2014 truly.\n", encoding="utf-8")
    test = check()["tests"]["prose"]
    assert test["state"] == "fail" and "U+2014" in test["why"] and "game changer" in test["why"]


def test_precheck_waits_on_the_close_and_on_a_person(acme):
    rules = acme["rules"] / bc.RULES_MD
    rules.write_text(RULES.replace("- Closed by hand: 2026-09\n", ""), encoding="utf-8")
    res = run("board_package_check", NAME, "--period", PERIOD, "--precheck")
    assert res.returncode == 0 and res.stdout.startswith("NOTHING: acme-board 2026-09: waiting on the close")
    rules.write_text(RULES, encoding="utf-8")
    res = run("board_package_check", NAME, "--period", PERIOD, "--precheck")
    assert res.stdout.startswith("NOTHING: acme-board 2026-09: waiting on the results")
    run("board_package_pack", NAME, "--period", PERIOD)
    res = run("board_package_check", NAME, "--period", PERIOD, "--precheck")
    assert res.stdout.startswith("WORK: acme-board 2026-09: the close is done and no package is drafted yet")
    (acme["package"] / "Acme Monthly Financial Report (2026-09).pdf").write_bytes(b"%PDF")
    res = run("board_package_check", NAME, "--period", PERIOD, "--precheck")
    assert res.stdout.startswith("NOTHING: acme-board 2026-09: a person already made the package")


def test_a_bad_argument_is_exit_2(acme):
    assert run("board_package_check").returncode == 2
    assert run("board_package_check", "nobody").returncode == 2
    assert run("board_package_check", NAME, "--period", "2026-13").returncode == 2
    res = run("board_package_check", NAME, "--period=", "--as-of=", "--precheck")
    assert res.returncode == 0 and res.stdout.startswith(("WORK: ", "NOTHING: "))
    assert run("board_package_tieout", NAME, "--period", PERIOD).returncode == 2  # no ledger yet


def test_rules_need_a_results_file_and_a_deck(acme):
    rules = acme["rules"] / bc.RULES_MD
    rules.write_text(RULES.replace("- Results file: `Acme {Month} {yyyy} Results*.html`\n", ""), encoding="utf-8")
    res = run("board_package_check", NAME, "--period", PERIOD)
    assert res.returncode == 2 and "Results file is required" in res.stderr
    rules.write_text(RULES.replace("DRAFT.pptx", "DRAFT.docx"), encoding="utf-8")
    assert run("board_package_check", NAME, "--period", PERIOD).returncode == 2


def test_the_evidence_file_keeps_other_rows_byte_for_byte(tmp_path):
    path = tmp_path / "ev.csv"
    path.write_bytes(b"\xef\xbb\xbf" + ",".join(bc.EVIDENCE_COLUMNS).encode() + b"\r\n"
                     + b'review:redteam,review,redteam,passed,,"a\r\nb",x,2026-10-01T00:00:00+00:00,Dana\r\n')
    before = path.read_bytes()
    assert bc.upsert_row(path, bc.EVIDENCE_COLUMNS, "review:finance", {"kind": "review", "state": "passed"}) == "created"
    after = path.read_bytes()
    assert after.startswith(before) and after.endswith(b"\r\n")
    assert bc.upsert_row(path, bc.EVIDENCE_COLUMNS, "review:finance", {"state": "failed"}) == "updated"
    assert bc.evidence_rows(path)["review:finance"]["state"] == "failed"
    assert path.read_bytes().startswith(before)
