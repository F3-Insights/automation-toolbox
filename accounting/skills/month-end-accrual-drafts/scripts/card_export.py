#!/usr/bin/env python3
# /// script
# dependencies = ["openpyxl"]
# ///
"""Turn a bank's commercial-card transaction download (xlsx) into the period's card CSV,
the file cc_accrual.py reads.

Keeps the rows whose transaction date falls in the period, drops payments to the card (they
are not expenses), and keeps the bank's signs, so a refund stays a credit. The cardholder is
the "diverted from" cardholder when the export has one (a charge moved onto a central account
still belongs to the person who made it), otherwise the cardholder column. The department a
person fills in per transaction is kept as dept_export.

The sheet and column names default to a common commercial-card detail export. Describe a
different export with --sheet and --column key=Header (repeatable); keys are date, post_date,
merchant, amount, cardholder, diverted, mcc, mcc_description, tran_type, dept_export,
reference. Headers match without regard to case or surrounding spaces. A missing required
column stops the run, naming it. When the sheet is not found, the first sheet is read.

Writes the CSV to --out (columns date, post_date, merchant, amount, cardholder, mcc,
mcc_description, tran_type, dept_export, reference). Prints a summary, or with --format json
period, row_count, total, dropped, passed and summary.
Exit 0 when rows were found, 1 when no row falls in the period, 2 on bad input.

Example:
    python3 card_export.py "August card transactions.xlsx" --period 2026-08 --out card-2026-08.csv
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

DEFAULT_SHEET = "CommercialCardTransactionDetail"
DEFAULT_COLUMNS = {
    "date": "Tran Date", "post_date": "Post Date", "merchant": "Merchant Name", "amount": "Amount",
    "cardholder": "Cardholder Name", "diverted": "Diverted From Cardholder Name", "mcc": "MCC Code",
    "mcc_description": "MCC Description", "tran_type": "Tran Type", "dept_export": "Dept",
    "reference": "Reference Number",
}
OPTIONAL = ("diverted", "dept_export", "post_date", "mcc_description", "reference")
FIELDS = ["date", "post_date", "merchant", "amount", "cardholder", "mcc", "mcc_description",
          "tran_type", "dept_export", "reference"]


def as_day(value) -> str:
    """A cell as YYYY-MM-DD, from a date or text such as '8/14/2026'."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if "/" in text:
        try:
            month, day, year = (int(p) for p in text.split()[0].split("/"))
            return f"{year:04d}-{month:02d}-{day:02d}"
        except ValueError:
            return text
    return text[:10]


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_grid(path: Path, sheet: str) -> list:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = book[sheet] if sheet in book.sheetnames else book[book.sheetnames[0]]
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        book.close()


def convert(grid: list, period: str, columns: dict | None = None) -> dict:
    """The period's card rows from a sheet grid whose first row is the header."""
    if not grid:
        raise ValueError("the sheet is empty")
    columns = {**DEFAULT_COLUMNS, **(columns or {})}
    at = {as_text(c).casefold(): i for i, c in enumerate(grid[0]) if as_text(c)}
    missing = [h for key, h in columns.items() if key not in OPTIONAL and h.casefold() not in at]
    if missing:
        raise ValueError(f"missing column(s) {missing}; the sheet has {[as_text(c) for c in grid[0] if c]}")

    def cell(row, key):
        index = at.get(columns[key].casefold())
        return row[index] if index is not None and index < len(row) else None

    kept, dropped = [], Counter()
    for row in grid[1:]:
        if cell(row, "date") is None:
            continue
        day, tran_type = as_day(cell(row, "date")), as_text(cell(row, "tran_type"))
        if not day.startswith(period):
            dropped[f"transaction date outside {period}"] += 1
            continue
        if tran_type.lower() == "payment":
            dropped["payment to the card (not an expense)"] += 1
            continue
        try:
            amount = float(cell(row, "amount") or 0)
        except (TypeError, ValueError):
            raise ValueError(f"amount {cell(row, 'amount')!r} on {day} is not a number") from None
        kept.append({
            "date": day, "post_date": as_day(cell(row, "post_date")),
            "merchant": as_text(cell(row, "merchant")), "amount": f"{amount:.2f}",
            "cardholder": (as_text(cell(row, "diverted")) or as_text(cell(row, "cardholder"))).upper(),
            "mcc": as_text(cell(row, "mcc")), "mcc_description": as_text(cell(row, "mcc_description")),
            "tran_type": tran_type, "dept_export": as_text(cell(row, "dept_export")).upper(),
            "reference": as_text(cell(row, "reference")),
        })
    total = round(sum(float(k["amount"]) for k in kept), 2)
    return {
        "period": period, "rows": kept, "row_count": len(kept), "total": total,
        "dropped": dict(dropped), "passed": bool(kept),
        "summary": f"{period}: {len(kept)} card row(s), ${total:,.2f}; dropped {dict(dropped) or 'none'}",
    }


def parse_columns(pairs: list) -> dict:
    out = {}
    for pair in pairs:
        key, sep, header = pair.partition("=")
        if not sep or key.strip() not in DEFAULT_COLUMNS:
            raise ValueError(f"--column {pair!r}: expected key=Header, key one of {sorted(DEFAULT_COLUMNS)}")
        out[key.strip()] = header.strip()
    return out


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0],
                                     formatter_class=argparse.RawDescriptionHelpFormatter,
                                     epilog=__doc__.split("\n\n", 1)[1])
    parser.add_argument("xlsx")
    parser.add_argument("--period", required=True, help="accounting period YYYY-MM")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="sheet name; the first sheet when absent")
    parser.add_argument("--column", action="append", default=[], help="key=Header; repeatable")
    parser.add_argument("--out", default="", help="write the card CSV here")
    parser.add_argument("--format", choices=["text", "json"], default="text")
    args = parser.parse_args(argv)
    try:
        result = convert(read_grid(Path(args.xlsx).expanduser(), args.sheet), args.period,
                         parse_columns(args.column))
        if args.out and result["rows"]:
            out = Path(args.out).expanduser()
            out.parent.mkdir(parents=True, exist_ok=True)
            with out.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=FIELDS)
                writer.writeheader()
                writer.writerows(result["rows"])
            result["out"] = str(out)
    except (OSError, ValueError, KeyError) as exc:
        print(f"card_export: {exc}", file=sys.stderr)
        return 2
    if args.format == "json":
        print(json.dumps({k: v for k, v in result.items() if k != "rows"}, indent=1))
    else:
        print(result["summary"])
        if result.get("out"):
            print(f"  wrote {result['out']}")
        print("PASS" if result["passed"] else f"FAIL: no card rows dated {args.period}")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
