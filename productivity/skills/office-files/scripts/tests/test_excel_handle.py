import csv
import json

import pytest
from openpyxl import Workbook

import excel_handle


@pytest.fixture
def workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Customer", "Amount", "Status"])
    ws.append(["Acme Components", 1200, "Open"])
    ws.append(["Northwind Traders", 450, "Paid"])
    ws.append(["Fabrikam Logistics", 3100, "Open"])
    ws.append([None, None, None])
    ws.append(["Lakeview Hardware", "=B2+B3", "Open"])
    wb.create_sheet("Notes").append(["Note"])
    path = tmp_path / "invoices.xlsx"
    wb.save(path)
    return str(path)


def run(capsys, *argv):
    assert excel_handle.main(list(argv)) == 0
    return json.loads(capsys.readouterr().out)


def test_analyze_lists_sheets_and_skips_blank_rows(workbook, capsys):
    out = run(capsys, "analyze", workbook)
    assert out["sheets"] == ["Invoices", "Notes"]
    assert out["rows"] == 4
    assert out["column_info"]["Amount"]["non_null_count"] == 3  # the formula has no cached value


def test_read_columns_and_rows(workbook, capsys, tmp_path):
    target = tmp_path / "out.csv"
    out = run(capsys, "read", workbook, "--columns", "Customer,Status", "--rows", "2",
              "--output", str(target))
    assert out["data"] == [{"Customer": "Acme Components", "Status": "Open"},
                           {"Customer": "Northwind Traders", "Status": "Paid"}]
    assert list(csv.reader(target.open()))[0] == ["Customer", "Status"]


def test_extract_numeric_and_contains(workbook, capsys):
    assert run(capsys, "extract", workbook, "--column", "Amount", "--condition", ">1000")["rows"] == 2
    out = run(capsys, "extract", workbook, "--column", "Customer", "--condition", "contains:north")
    assert out["unique_values"] == ["Northwind Traders"]
    assert run(capsys, "extract", workbook, "--column", "Status", "--condition", "==Open")["rows"] == 3


def test_convert_all_sheets(workbook, capsys, tmp_path):
    out = run(capsys, "convert", workbook, "--all-sheets", "--output", str(tmp_path / "x.csv"))
    assert [p.rsplit("/", 1)[-1] for p in out["converted_files"]] == ["x_Invoices.csv", "x_Notes.csv"]


def test_process_stats_and_frequency(workbook, capsys):
    stats = run(capsys, "process", workbook, "--column", "Amount", "--col-operation", "stats")
    assert stats["max"] == 3100 and stats["min"] == 450
    freq = run(capsys, "process", workbook, "--column", "Status", "--col-operation", "frequency")
    assert freq["frequency"] == {"Open": 3, "Paid": 1}


def test_missing_column_fails(workbook):
    with pytest.raises(SystemExit) as exit_info:
        excel_handle.main(["extract", workbook, "--column", "Vendor"])
    assert exit_info.value.code == 1


def test_repeated_header_keeps_both_columns(tmp_path, capsys):
    wb = Workbook()
    wb.active.append(["Account", "Amount", "Amount"])
    wb.active.append(["4000", 100, 250])
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    out = run(capsys, "read", str(path))
    assert out["columns"] == ["Account", "Amount", "Amount.1"]
    assert out["data"][0]["Amount"] == 100 and out["data"][0]["Amount.1"] == 250


@pytest.fixture
def big_workbook(tmp_path):
    wb = Workbook()
    wb.active.append(["Line", "Amount"])
    for n in range(1, 451):
        wb.active.append([n, n * 10])
    path = tmp_path / "big.xlsx"
    wb.save(path)
    return str(path)


def test_read_and_extract_print_at_most_200_rows_by_default(big_workbook, capsys):
    for out in (run(capsys, "read", big_workbook),
                run(capsys, "extract", big_workbook, "--column", "Amount", "--condition", ">0")):
        assert len(out["data"]) == 200 and out["rows"] == 450 and out["rows_shown"] == 200
        assert "first 200 of 450" in out["note"] and "--limit 0" in out["note"]
        assert out["data"][-1]["Line"] == 200


def test_limit_n_and_limit_zero(big_workbook, capsys):
    out = run(capsys, "read", big_workbook, "--limit", "5")
    assert len(out["data"]) == 5 and "first 5 of 450" in out["note"]
    out = run(capsys, "read", big_workbook, "--limit", "0")
    assert len(out["data"]) == 450 and "note" not in out and "rows_shown" not in out


def test_no_note_when_nothing_was_cut(workbook, capsys):
    out = run(capsys, "read", workbook)
    assert len(out["data"]) == 4 and "note" not in out


def test_output_file_keeps_every_row_while_stdout_is_limited(big_workbook, capsys, tmp_path):
    target = tmp_path / "all.csv"
    out = run(capsys, "read", big_workbook, "--limit", "10", "--output", str(target))
    assert len(out["data"]) == 10
    assert len(list(csv.reader(target.open()))) == 451
    as_json = tmp_path / "all.json"
    run(capsys, "extract", big_workbook, "--column", "Line", "--output", str(as_json))
    assert len(json.loads(as_json.read_text())["data"]) == 450


def test_negative_limit_is_refused(big_workbook):
    with pytest.raises(SystemExit) as exit_info:
        excel_handle.main(["read", big_workbook, "--limit", "-1"])
    assert exit_info.value.code == 2
