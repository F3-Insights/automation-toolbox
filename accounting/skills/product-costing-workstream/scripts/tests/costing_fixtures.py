"""Invented cost-model workbooks and layout for the product-costing tests."""

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FIELDS = ["sales_qty", "sales_amount", "prod_qty", "std_dm", "var_dm", "std_dl", "var_dl",
          "std_voh", "var_voh", "std_foh", "var_foh", "direct_cost", "total_cost", "cm", "gp", "cm_pct", "gp_pct"]
LAYOUT = {
    "month_tab": "Margin by month", "ytd_tab": "Margin YTD", "tb_tab": "Ledger", "alloc_tab": "Allocation",
    "first_row": 3, "info_columns": {"customer_group": 0, "customer": 1, "product_fg": 2, "description": 3},
    "block_start": 5, "block_width": 18, "block_fields": FIELDS, "periods": 12,
    "tb_columns": {"category": 0, "subcategory": 1, "description": 2}, "tb_first_amount_col": 3,
    "revenue_category": "Sales",
}


def write_layout(tmp_path, **changes):
    path = tmp_path / "layout.json"
    path.write_text(json.dumps({**LAYOUT, **changes}), encoding="utf-8")
    return path


def green_row(info, blocks):
    row = list(info) + [None] * (LAYOUT["block_start"] - len(info))
    for block in blocks:
        cells = [block.get(f) for f in FIELDS]
        row += cells + [None] * (LAYOUT["block_width"] - len(cells))
    return row


def tb_row(category, amounts, sub=None):
    return [category, sub, None] + list(amounts) + [0] * (12 - len(amounts))


def build(path, green, tb, ytd=None, alloc=None, extra_tabs=None):
    """green/ytd: list of (info, [block per month]); tb: list of TB rows; alloc: {cell: formula}."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = LAYOUT["month_tab"]
    tabs = [(ws, [green_row(i, b) for i, b in green]), (wb.create_sheet(LAYOUT["tb_tab"]), tb)]
    if ytd:
        tabs.append((wb.create_sheet(LAYOUT["ytd_tab"]), [green_row(i, b) for i, b in ytd]))
    for name, rows in (extra_tabs or {}).items():
        tabs.append((wb.create_sheet(name), rows))
    for sheet, rows in tabs:
        for _ in range(LAYOUT["first_row"] - 1):
            sheet.append([])
        for r in rows:
            sheet.append(r)
    if alloc:
        a = wb.create_sheet(LAYOUT["alloc_tab"])
        for coord, formula in alloc.items():
            a[coord] = formula
    wb.save(path)
    return path


def block(sales, cost, **kw):
    b = {"sales_qty": 10, "sales_amount": sales, "total_cost": cost, "cm": sales - cost,
         "cm_pct": (sales - cost) / sales if sales and cost else None}
    b.update(kw)
    return b
