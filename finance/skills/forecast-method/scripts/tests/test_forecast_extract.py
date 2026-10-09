from datetime import datetime

import pytest
import yaml

import forecast_extract
from _common import extract, month_of_header, section_of
from conftest import PRIOR_EBITDA, amounts, write_book


def test_extract_reads_the_consolidated_block_and_ties(folder):
    x = extract(folder / "delivery" / "Northwind forecast rev4.xlsx", folder)
    assert x["ok"], [c for c in x["checks"] if not c["ok"]]
    assert len(x["rows"]) == 11                                  # the second entity is not counted
    t = x["years"]["2027"]["totals"]
    assert t["revenue"] == 12 * 93500
    assert t["cogs"] == 12 * 36000                               # the warehouse bonus is its own line
    assert t["bonus"] == 12 * 2300
    assert t["ebitda_pre_bonus"] == PRIOR_EBITDA
    assert t["ebitda"] == PRIOR_EBITDA - 12 * 2300
    assert sum(x["years"]["2027"]["quarters"].values()) == pytest.approx(PRIOR_EBITDA)


def test_a_broken_tie_exits_1(folder, tmp_path, capsys):
    bad = write_book(tmp_path / "bad.xlsx", amounts(4), break_tie=True)
    assert forecast_extract.main([str(bad), "--folder", str(folder)]) == 1
    assert "tie ebitda" in capsys.readouterr().out


def test_a_missing_column_exits_2(folder, capsys):
    settings = yaml.safe_load((folder / "FORECAST-SETTINGS.yaml").read_text())
    settings["layouts"]["default"]["columns"]["account"] = "ACCOUNT_NO"
    (folder / "FORECAST-SETTINGS.yaml").write_text(yaml.safe_dump(settings))
    assert forecast_extract.main([str(folder / "delivery" / "Northwind forecast rev4.xlsx"), "--folder", str(folder)]) == 2
    assert "ACCOUNT_NO" in capsys.readouterr().err


@pytest.mark.parametrize("cell,month", [
    (datetime(2027, 3, 1), "2027-03"), ("Month Ended March 2027", "2027-03"), ("2027-03", "2027-03"),
    ("Mar 2027", "2027-03"), ("Mar'27", "2027-03"), ("FY2027", ""), ("Total", ""), (None, "")])
def test_months_are_located_by_their_header(cell, month):
    assert month_of_header(cell) == month


def test_sections_take_the_longest_prefix():
    assert section_of("66900", {"6": "sga", "66900": "cogs"}) == "cogs"
    assert section_of("66100", {"6": "sga", "66900": "cogs"}) == "sga"
    assert section_of("15000", {"4": "revenue"}) == "other"


def test_a_month_headed_twice_fails_a_control(folder, tmp_path):
    from openpyxl import load_workbook
    book = load_workbook(folder / "delivery" / "Northwind forecast rev4.xlsx")
    book["Income Statement"].cell(3, 4 + 8).value = "Month Ended May 2027"   # September's column
    broken = tmp_path / "twice.xlsx"
    book.save(broken)
    x = extract(broken, folder)
    failed = {c["name"]: c["detail"] for c in x["checks"] if not c["ok"]}
    assert "2027-05 (columns 8 and 12)" in failed["every month column is headed once"]
    assert "2027-09" in failed["no month is missing between the first and the last"]


def test_a_band_keeps_only_the_named_column_blocks(folder, tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = "IS"
    ws.append([None, None, "Northwind Traders"])
    ws.append([None, None, None] + ["Plan"] * 12 + ["Actual"] * 4 + ["Forecast"] * 8)
    ws.append([None, None, None] + [datetime(2027, m, 1) for m in range(1, 13)] + [f"{m}'27" for m in
              ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")])
    ws.append([41000, "Wholesale revenue", None] + [555.0] * 12 + [200.0] * 12)
    ws.append([62000, "Lease", None] + [555.0] * 12 + [50.0] * 12)
    ws.append([None, "Total Revenue", None] + [555.0] * 12 + [200.0] * 12)
    path = tmp_path / "is.xlsx"
    book.save(path)
    settings = yaml.safe_load((folder / "FORECAST-SETTINGS.yaml").read_text())
    settings["layouts"]["is"] = {"sheet": "IS", "header_row": 3, "band": {"row": 2, "include": ["Actual", "Forecast"]},
                                 "columns": {"account": "col:A", "account_name": "col:B"}, "blank_token": "",
                                 "tie_rows": {"revenue": "Total Revenue"}}
    (folder / "FORECAST-SETTINGS.yaml").write_text(yaml.safe_dump(settings))
    x = extract(path, folder, "is")
    assert x["ok"], [c for c in x["checks"] if not c["ok"]]
    assert x["years"]["2027"]["totals"]["revenue"] == 2400.0
    assert x["years"]["2027"]["totals"]["ebitda"] == 1800.0
