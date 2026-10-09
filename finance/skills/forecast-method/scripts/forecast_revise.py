# /// script
# dependencies = ["pyyaml", "openpyxl"]
# ///
"""Build a rolling vintage's workbook as a values-only copy of the prior forecast.

Only for a vintage of kind `rolling` (the monthly reforecast the agents build), never for a
revision a person issued. Reads the vintage's proposals.json (one row per account x department x
location x class, the amount for each forecast month in income-statement sign, and the `node` or
reason that put it there), then:

1. copies the prior workbook to the vintage's new_workbook path (the prior is only read);
2. writes the proposed amounts as values into the sheet the settings' `revise: {layout: ...}`
   names (identity columns plus one column per month); a row the sheet lacks is appended, a
   contra-revenue account is written back in the sheet's own sign, a month at or before
   last_actual_month is refused, and a formula cell is never overwritten;
3. writes change-log.csv (sheet, cell, old, new, node, at) and upload.csv (the sheet's own
   columns, values only, for a person to upload to the accounting system);
4. reopens the copy and checks the package: content types and workbook relationships parse,
   every sheet target exists, every sheet the prior had is still there.

A prior workbook with dynamic-array formulas (xl/metadata.xml) is refused: a library save
damages them, so a person builds that revision.

Exit 0 when the copy is written and passes its checks, 1 when an integrity check fails, 2 on a
refusal or a bad argument. --dry-run says what would change and writes nothing.

Example:
    python3 forecast_revise.py ~/Forecast --vintage 2026-10
"""

import argparse
import csv
import json
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from _common import (ForecastError, cents, close, forecast_folder, ident, layout, load_settings, load_vintage,
                     month_of_header, now, resolve_path, row_key)

REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def has_dynamic_arrays(path):
    with zipfile.ZipFile(path) as archive:
        return "xl/metadata.xml" in archive.namelist()


def package_parts(path):
    """The package's plumbing resolves: what Excel's repair dialog would otherwise tell a person."""
    out = {"content_types_parse": False, "workbook_rels_parse": False, "sheet_targets_exist": False}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        try:
            ElementTree.fromstring(archive.read("[Content_Types].xml"))
            out["content_types_parse"] = True
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            out["workbook_rels_parse"] = True
        except (KeyError, ElementTree.ParseError):
            return out
        targets = [e.get("Target", "") for e in rels.iter(REL) if str(e.get("Type", "")).endswith("/worksheet")]
        resolved = [t[1:] if t.startswith("/") else (t if t.startswith("xl/") else f"xl/{t}") for t in targets]
        out["sheet_targets_exist"] = bool(resolved) and all(t in names for t in resolved)
    return out


def write_upload(workbook, sheet, header_row, path):
    """The sheet as the accounting system's upload: its own columns, values only, formulas blank."""
    from openpyxl import load_workbook
    book = load_workbook(str(workbook), read_only=True)
    try:
        rows = list(book[sheet].iter_rows(min_row=header_row, values_only=True))
    finally:
        book.close()
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in rows:
            cells = ["" if v is None or is_formula(v) else (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v)
                     for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                writer.writerow(cells)


def revise(folder, vintage_id, dry_run=False):
    root = forecast_folder(folder)
    settings = load_settings(root)
    meta = load_vintage(root, vintage_id)
    if meta.get("kind") != "rolling":
        raise ForecastError(f"vintage {vintage_id} is a {meta.get('kind')}: only a rolling vintage is built here")
    spec_name = str((settings.get("revise") or {}).get("layout") or "")
    if not spec_name:
        raise ForecastError("settings need `revise: {layout: <name>}`, the upload-shaped sheet a rolling vintage writes")
    spec = layout(settings, spec_name)
    vdir = Path(meta["_dir"])
    prior = resolve_path(root, meta["prior_workbook"])
    target = resolve_path(root, meta["new_workbook"])
    if target == prior:
        raise ForecastError("the new workbook path is the prior workbook: a person's file is never written")
    if not (vdir / "proposals.json").is_file():
        raise ForecastError("proposals.json is not in the vintage folder")
    if has_dynamic_arrays(prior):
        raise ForecastError(f"{prior.name} carries dynamic-array formulas, which a library save damages; "
                            "a person builds this revision")
    last = meta["last_actual_month"]
    if not last:
        raise ForecastError("a rolling vintage needs last_actual_month in VINTAGE.yaml")
    proposals = json.loads((vdir / "proposals.json").read_text(encoding="utf-8")).get("rows") or []
    for row in proposals:
        early = [m for m in row.get("amounts") or {} if m <= last]
        if early:
            raise ForecastError(f"{row.get('account')}: {', '.join(early)} is closed; actuals are never proposed")
    blank = str(spec.get("blank_token", "blank") or "")
    contra = {str(a) for a in spec.get("contra_revenue_accounts") or []}

    from openpyxl import load_workbook
    work = target.with_name(target.stem + ".tmp.xlsx")
    if not dry_run:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(prior, work)
    book = load_workbook(str(prior if dry_run else work))
    sheet_name = str(spec["sheet"])
    if sheet_name not in book.sheetnames:
        raise ForecastError(f"{prior.name} has no sheet '{sheet_name}'")
    original = list(book.sheetnames)
    sheet = book[sheet_name]
    header_row = int(spec["header_row"])
    header = [c.value for c in sheet[header_row]]
    labels = ["" if h is None else str(h).strip().lower() for h in header]
    columns = {}
    for name, text in (spec.get("columns") or {}).items():
        if str(text or "").strip():
            if str(text).strip().lower() not in labels:
                raise ForecastError(f"{sheet_name}: no column headed '{text}'")
            columns[name] = labels.index(str(text).strip().lower())
    months = {}
    for position, cell in enumerate(header):
        months.setdefault(month_of_header(cell), position)
    months.pop("", None)

    def cell_text(values, name):
        return ident(values[columns[name]], blank) if name in columns and columns[name] < len(values) else ""

    existing = {}
    for index in range(header_row + 1, sheet.max_row + 1):
        values = [c.value for c in sheet[index]]
        ids = {name: cell_text(values, name) for name in ("account", "department", "location", "cls")}
        if ids["account"]:
            existing.setdefault(row_key(ids), index)

    stamp, changes, skipped = now(), [], []
    next_row = sheet.max_row + 1
    for row in proposals:
        ids = {k: str(row.get(k) or "") for k in ("account", "department", "location", "cls")}
        index = existing.get(row_key(ids))
        if index is None:   # a row the sheet lacks: append it with its identity cells
            index = existing[row_key(ids)] = next_row
            next_row += 1
            for name, at in columns.items():
                value = str(row.get("account_name") or "") if name == "account_name" else ids.get(name, "")
                sheet.cell(index, at + 1).value = value or (blank if name in ("department", "cls", "location") else "")
        flip = -1.0 if ids["account"] in contra else 1.0
        for month, amount in sorted((row.get("amounts") or {}).items()):
            if month not in months:
                raise ForecastError(f"{sheet_name}: no column for {month}")
            cell = sheet.cell(index, months[month] + 1)
            old = cell.value
            if is_formula(old):
                skipped.append(f"{cell.coordinate} holds a formula")
                continue
            wanted = cents(flip * float(amount))
            if isinstance(old, (int, float)) and close(float(old), wanted, float(settings["tolerance"])):
                continue
            cell.value = wanted
            changes.append([sheet_name, cell.coordinate, "" if old is None else str(old), f"{wanted:.2f}",
                            str(row.get("node") or ""), stamp])
    result = {"vintage": vintage_id, "target": str(target), "changes": len(changes), "skipped": skipped,
              "dry_run": dry_run}
    if dry_run:
        book.close()
        return dict(result, ok=True)
    book.save(str(work))
    book.close()
    work.replace(target)
    integrity = package_parts(target)
    check = load_workbook(str(target), read_only=True)
    integrity["sheets_present"] = all(name in check.sheetnames for name in original)
    check.close()
    with (vdir / "change-log.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["sheet", "cell", "old", "new", "node", "at"])
        writer.writerows(changes)
    write_upload(target, sheet_name, header_row, vdir / "upload.csv")
    return dict(result, integrity=integrity, ok=all(integrity.values()),
                files=[str(target), str(vdir / "change-log.csv"), str(vdir / "upload.csv")])


def main(argv=None):
    parser = argparse.ArgumentParser(description="Build a rolling vintage's workbook as a values-only copy of the prior.")
    parser.add_argument("folder", help="the Forecast folder")
    parser.add_argument("--vintage", required=True, help="a rolling vintage under vintages/")
    parser.add_argument("--dry-run", action="store_true", help="say what would change; write nothing")
    parser.add_argument("--format", choices=["text", "json"], default="text")
    args = parser.parse_args(argv)
    try:
        result = revise(args.folder, args.vintage, args.dry_run)
    except (ForecastError, KeyError) as exc:
        print(f"forecast-revise: {exc}", file=sys.stderr)
        return 2
    if args.format == "json":
        print(json.dumps(result, indent=1))
    else:
        head = "DRY RUN" if args.dry_run else ("OK" if result["ok"] else "FAIL")
        print(f"{head}: {result['changes']} cell(s) to {Path(result['target']).name}"
              + (f"; {len(result['skipped'])} formula cell(s) left alone" if result["skipped"] else ""))
        for name, ok in (result.get("integrity") or {}).items():
            if not ok:
                print(f"  FAIL integrity {name}")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
