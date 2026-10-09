"""What the board-package scripts share: the company, its rules, the period and its files,
how a printed number is read, where figures are found in a file, the tie-out and the
evidence file.

THE COMPANY
    COMPANY is a Context name or the path to a Context YAML file. A name is looked up as
    <contexts>/<name>.yaml, where <contexts> is --contexts-dir, else the owner setting
    contexts_dir in [board-package-workstream], else the top-level contexts_dir.
    The Context lists sources ({name, path}); the scripts use three folders by name:
    rules (holds BOARD-RULES.md), close (the Month-End folder, {yyyy}/{yyyy-mm}/ per month,
    read only) and package (the Reporting folder). Sources named prior or prior-* are more
    places earlier packages are found.

THE RULES
    BOARD-RULES.md holds "- Key: value" lines under "## Package inputs": Results file
    (required, a glob in the close's month folder), Results unit (1, k or M), Period folder,
    Deck file (required, .pptx or .html), Script file, Cover email, Lender pack, Review note,
    Deck template, Previous packages, Supporting files, Close gate (month-end-check or
    attested), Closed by hand, Tie-out tolerance, Banned phrases, Banned characters, Figure
    ignore. Patterns take {yyyy}, {mm}, {yyyy-mm}, {Month} and {Mon}.

THE PERIOD FOLDER
    <package>/<Period folder>, or --period-dir. Its work/ folder holds sources.json,
    sources.md, results-figures.csv, package-figures.csv, BOARD-EVIDENCE-<p>.csv,
    tieout-<p>.json, finance-review.md, redteam.md, CONFIRMATIONS.md and LOG.md. The tie-out
    ledger, tieout-ledger-<p>.csv, sits beside the package files.
"""

from __future__ import annotations

import ast
import calendar
import csv
import glob
import io
import json
import operator
import os
import re
import subprocess
import sys
import tempfile
import tomllib
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from html.parser import HTMLParser
from pathlib import Path

RULES_MD = "BOARD-RULES.md"
SOURCES_JSON = "sources.json"
SOURCES_MD = "sources.md"
RESULTS_FIGURES_CSV = "results-figures.csv"
PACKAGE_FIGURES_CSV = "package-figures.csv"
FINANCE_REVIEW_MD = "finance-review.md"
REDTEAM_MD = "redteam.md"
CONFIRMATIONS_MD = "CONFIRMATIONS.md"
SOURCES_SCHEMA = "board-package/sources/1"

RESULTS_COLUMNS = ("id", "file", "period", "metric", "value", "unit", "printed", "location")
PACKAGE_COLUMNS = ("id", "file", "location", "printed", "unit", "metric", "source", "note")
EVIDENCE_COLUMNS = ("id", "kind", "item", "state", "evidence", "review_file", "note", "updated_at", "by")
REVIEW_STATES = ("passed", "failed")
EXCEPTION_STATES = ("approved", "withdrawn")
REVIEW_KINDS = ("finance", "redteam")

UNITS = {"": 1.0, "1": 1.0, "$": 1.0, "k": 1e3, "m": 1e6, "mm": 1e6, "b": 1e9, "bn": 1e9, "%": 1.0}
PERIOD_RE = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
VERSION_RE = re.compile(r"\bv(\d+)\b", re.I)
CONFIRM_HEAD_RE = re.compile(r"^## (CR-[0-9a-f]{10})\s*$", re.M)
HOUSE_PHRASES = ("on track", "ahead of schedule", "great progress", "successfully completed",
                 "key milestone", "critical finding", "alarming", "major concern", "game-changer",
                 "revolutionary", "seamless", "robust", "leverage", "going forward",
                 "it should be noted", "as previously discussed", "circle back", "cutting-edge")
MONTH_END_CHECK = "~/.claude/skills/month-end-workstream/scripts/month_end_check.py"


class Bad(Exception):
    """A bad argument, a missing Context, source or rules file: exit 2."""


class Refused(Exception):
    """A write refused because it would break the contract: exit 1."""


def run_main(main) -> None:
    """Run a script's main(), turning Refused into exit 1 and Bad into exit 2."""
    try:
        code = main()
    except Refused as exc:
        print(f"REFUSED {exc}", file=sys.stderr)
        sys.exit(1)
    except Bad as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        sys.exit(2)
    sys.exit(code or 0)


def settings(section=None):
    """Owner settings from ~/.config/f3i-toolbox/settings.toml (or $F3I_TOOLBOX_SETTINGS)."""
    path = Path(os.environ.get("F3I_TOOLBOX_SETTINGS", "~/.config/f3i-toolbox/settings.toml")).expanduser()
    try:
        data = tomllib.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        data = {}
    return data.get(section, {}) if section else data


def blank(value) -> str | None:
    """A form left blank sends --x=; an empty or whitespace value counts as not given."""
    text = "" if value is None else str(value).strip()
    return text or None


def now_utc() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def stamp() -> str:
    return datetime.now().strftime("%Y%m%dT%H%M%S")


def split_list(value, sep=","):
    return [p.strip().strip("`").strip() for p in (value or "").split(sep) if p.strip().strip("`").strip()]


def add_period_options(parser) -> None:
    """The options every per-company script takes."""
    parser.add_argument("--period", default="", help="yyyy-mm; blank: the month before today")
    parser.add_argument("--period-dir", default="", help="Use this folder as the period folder (a dry run's)")
    parser.add_argument("--as-of", default="", help="Today, for a blank period (yyyy-mm-dd)")
    parser.add_argument("--contexts-dir", default="", help="Where Context YAML files are found by name")


# --------------------------------------------------------------------------- printed numbers

NUM = r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?|\.\d+"
TOKEN_RE = re.compile(
    r"(?<![A-Za-z0-9_.,/:])"
    r"(?P<sign>[+\-\u2212\u2013])?(?P<cur1>\$)?(?P<open>\()?(?P<sign2>[+\-\u2212])?(?P<cur2>\$)?"
    rf"(?P<num>{NUM})"
    r"(?P<close>\))?(?:\s?(?P<suf>[kK]|MM|mm|M|bn|B)(?![A-Za-z]))?(?P<close2>\))?(?P<pct>%)?"
    r"(?![0-9])")
DATE_LIKE = re.compile(r"\b\d{4}-\d{2}(?:-\d{2})?\b|\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b|\b\d{1,2}:\d{2}\b"
                       r"|\b\d{1,2}(?:st|nd|rd|th)\b")


@dataclass
class Printed:
    """One number as printed: its value in printed units, the decimals shown, its suffix
    (k, m, b), whether it is a percentage, and whether it carries a mark that makes it a
    figure ($, a suffix, %, parentheses, a thousands comma, a decimal point, a sign)."""

    text: str
    value: float
    decimals: int
    suffix: str
    percent: bool
    marked: bool

    @property
    def multiplier(self) -> float:
        return UNITS.get(self.suffix, 1.0)


def _parse_match(m) -> Printed:
    text = m.group(0).strip()
    raw = m.group("num")
    opened, closed = bool(m.group("open")), bool(m.group("close") or m.group("close2"))
    if opened != closed:
        # "(74" or "74)": the bracket belongs to the sentence, not the number.
        text = text.replace("(", "", 1) if opened else text.rstrip(")")
        opened = closed = False
    sign = (m.group("sign") or "") + (m.group("sign2") or "")
    if opened and closed and "+" in (m.group("sign2") or ""):
        # "(+$104k)": a positive figure in a parenthetical, not a negative one.
        text = text.strip("()")
        opened = closed = False
    value = float(raw.replace(",", ""))
    if (opened and closed) or any(c in sign for c in "-\u2212\u2013"):
        value = -value
    suffix = (m.group("suf") or "").lower()
    suffix = {"mm": "m", "bn": "b"}.get(suffix, suffix)
    percent = bool(m.group("pct"))
    marked = bool(m.group("cur1") or m.group("cur2") or suffix or percent or (opened and closed)
                  or "," in raw or "." in raw or sign)
    decimals = len(raw.split(".", 1)[1]) if "." in raw else 0
    return Printed(text, value, decimals, suffix, percent, marked)


def parse_printed(text) -> Printed | None:
    """The one number a cell or a ledger's printed value holds, or None."""
    s = (text or "").strip().replace("\u00a0", " ").replace("\u202f", " ")
    if not s:
        return None
    m = TOKEN_RE.fullmatch(s) or TOKEN_RE.fullmatch(s.replace(" ", ""))
    return _parse_match(m) if m else None


def tokens(text, ignore=()) -> list[Printed]:
    """Every number in running text; dates and times are blanked out first."""
    s = (text or "").replace("\u00a0", " ").replace("\u202f", " ")
    s = DATE_LIKE.sub(lambda m: " " * len(m.group(0)), s)
    found = [_parse_match(m) for m in TOKEN_RE.finditer(s)]
    return [p for p in found if not any(rx.search(p.text) for rx in ignore)]


def is_year(p: Printed) -> bool:
    return (not p.marked) and p.decimals == 0 and 1900 <= abs(p.value) <= 2100


def figure_in_text(p: Printed) -> bool:
    """In running text only a marked number is a figure: '3 sites' and 'Entity 11' are not."""
    return p.marked and not is_year(p)


def unit_multiplier(unit) -> float:
    key = (blank(unit) or "").lower().lstrip("$")
    key = {"thousands": "k", "000": "k", "000s": "k", "millions": "m", "mm": "m"}.get(key, key)
    if key not in UNITS:
        raise Bad(f"unit {unit!r} is not one of 1, k, M, B or %")
    return UNITS[key]


def dollars(p: Printed, unit) -> float:
    """The printed value in base units: its own suffix wins, else the given unit."""
    if p.percent:
        return p.value
    return p.value * (p.multiplier if p.suffix else unit_multiplier(unit))


def round_to_print(value: float, p: Printed, unit) -> float:
    """value (base units) rounded half up to the precision p is printed at."""
    mult = 1.0 if p.percent else (p.multiplier if p.suffix else unit_multiplier(unit))
    scaled = Decimal(repr(value / mult)).quantize(Decimal(1).scaleb(-p.decimals), rounding=ROUND_HALF_UP)
    return float(scaled) * mult


# --------------------------------------------------------------------------- periods

def check_period(period) -> str:
    if not PERIOD_RE.match(period or ""):
        raise Bad(f"period {period!r} is not yyyy-mm")
    return period


def prior_period(period: str) -> str:
    year, month = int(period[:4]), int(period[5:])
    return f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"


def resolve_period(period, as_of=None) -> str:
    if blank(period):
        return check_period(blank(period))
    try:
        today = date.fromisoformat(blank(as_of)) if blank(as_of) else date.today()
    except ValueError:
        raise Bad(f"--as-of {as_of!r} is not yyyy-mm-dd") from None
    return prior_period(f"{today.year}-{today.month:02d}")


def fill(pattern: str, period: str) -> str:
    name = calendar.month_name[int(period[5:])]
    return (pattern.replace("{yyyy-mm}", period).replace("{yyyy}", period[:4]).replace("{mm}", period[5:])
            .replace("{Month}", name).replace("{Mon}", name[:3]))


# --------------------------------------------------------------------------- the company and its rules

def context_file(company, contexts=None) -> Path:
    name = blank(company)
    if not name:
        raise Bad("COMPANY is required: a Context name or the path to a Context YAML file")
    as_path = Path(name).expanduser()
    if as_path.suffix in (".yaml", ".yml") or "/" in name or os.sep in name:
        if not as_path.is_file():
            raise Bad(f"no Context file at {as_path}")
        return as_path
    folder = (blank(contexts) or blank(settings("board-package-workstream").get("contexts_dir"))
              or blank(settings().get("contexts_dir")))
    if not folder:
        raise Bad("set contexts_dir (in [board-package-workstream] or at the top level of the owner "
                  "settings), pass --contexts-dir, or give the Context file's path")
    for suffix in (".yaml", ".yml"):
        candidate = Path(folder).expanduser() / f"{name}{suffix}"
        if candidate.is_file():
            return candidate
    raise Bad(f"no Context named {name!r} in {folder} (pass --contexts-dir or a path)")


def field_lines(text: str, heading: str) -> dict:
    """The '- Key: value' lines under the heading, keys lower-cased, backticks unwrapped."""
    out, inside = {}, False
    for line in text.splitlines():
        found = re.match(r"^\s*#{1,6}\s+(.*?)\s*#*\s*$", line)
        if found:
            inside = found.group(1).strip().lower() == heading
            continue
        found = re.match(r"^\s*[-*]\s+([^:]+?)\s*:\s*(.*?)\s*$", line) if inside else None
        if found:
            key = re.sub(r"\s+", " ", found.group(1).strip().strip("*_").strip()).lower()
            value = found.group(2).strip()
            if len(value) >= 2 and value[0] == value[-1] == "`":
                value = value[1:-1].strip()
            out[key] = value
    return out


@dataclass
class Rules:
    results_file: str
    results_unit: str
    period_folder: str
    deck_file: str
    script_file: str | None
    cover_email: str | None
    lender_pack: str | None
    review_note: str
    deck_template: str | None
    previous_packages: list
    supporting_files: list
    close_gate: str
    closed_by_hand: list
    tolerance: float
    banned_phrases: list
    banned_characters: list
    figure_ignore: list = field(default_factory=list)


FIGURE_SUFFIXES = (".pptx", ".html", ".htm", ".md", ".txt", ".xlsx")


def parse_rules(text: str) -> Rules:
    raw = field_lines(text, "package inputs")
    results, deck = blank(raw.get("results file")), blank(raw.get("deck file"))
    if not results:
        raise Bad(f"{RULES_MD}: Results file is required under '## Package inputs'")
    if not deck:
        raise Bad(f"{RULES_MD}: Deck file is required")
    if not deck.lower().endswith((".pptx", ".html")):
        raise Bad(f"{RULES_MD}: the Deck file must end .pptx or .html, not {deck!r}")
    for key in ("script file", "cover email"):
        value = blank(raw.get(key))
        if value and not value.lower().endswith((".md", ".txt")):
            raise Bad(f"{RULES_MD}: the {key.capitalize()} must end .md or .txt, not {value!r}")
    lender = blank(raw.get("lender pack"))
    if lender and lender.lower() in ("none", "no", "off"):
        lender = None
    if lender and not lender.lower().endswith(FIGURE_SUFFIXES):
        raise Bad(f"{RULES_MD}: the Lender pack must end {', '.join(FIGURE_SUFFIXES)}, not {lender!r}")
    unit = blank(raw.get("results unit")) or "1"
    unit_multiplier(unit)
    gate = (blank(raw.get("close gate")) or "month-end-check").lower()
    if gate not in ("month-end-check", "attested"):
        raise Bad(f"{RULES_MD}: Close gate must be month-end-check or attested, not {gate!r}")
    closed = split_list(raw.get("closed by hand"))
    for period in closed:
        if not PERIOD_RE.match(period):
            raise Bad(f"{RULES_MD}: Closed by hand {period!r} is not yyyy-mm")
    try:
        tolerance = float(blank(raw.get("tie-out tolerance")) or 1)
    except ValueError:
        raise Bad(f"{RULES_MD}: Tie-out tolerance {raw.get('tie-out tolerance')!r} is not a number") from None
    ignore = []
    for pattern in split_list(raw.get("figure ignore"), ";"):
        try:
            ignore.append(re.compile(pattern))
        except re.error as exc:
            raise Bad(f"{RULES_MD}: Figure ignore {pattern!r} is not a regular expression: {exc}") from None
    return Rules(
        results_file=results, results_unit=unit, period_folder=blank(raw.get("period folder")) or "{yyyy-mm}",
        deck_file=deck, script_file=blank(raw.get("script file")), cover_email=blank(raw.get("cover email")),
        lender_pack=lender, review_note=blank(raw.get("review note")) or "{yyyy-mm} package review notes.md",
        deck_template=blank(raw.get("deck template")),
        previous_packages=split_list(raw.get("previous packages"), ";"),
        supporting_files=split_list(raw.get("supporting files"), ";"),
        close_gate=gate, closed_by_hand=closed, tolerance=tolerance,
        banned_phrases=split_list(raw.get("banned phrases"), ";"),
        banned_characters=(blank(raw.get("banned characters")) or "").split() or ["\u2014"],
        figure_ignore=ignore)


@dataclass
class Company:
    name: str
    rules_file: Path
    rules: Rules
    close: Path
    package: Path
    prior: list


def load_company(company, contexts=None) -> Company:
    import yaml  # only needed here, so --help works without it
    path = context_file(company, contexts)
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise Bad(f"{path} could not be read as YAML: {type(exc).__name__}") from None
    if not isinstance(data, dict):
        raise Bad(f"{path} is not a Context: the top level is not a mapping")
    folders, prior = {}, []
    for source in data.get("sources") or []:
        if not isinstance(source, dict):
            continue
        name, where = str(source.get("name") or "").strip(), str(source.get("path") or "").strip()
        if not (name and where):
            continue
        if name == "prior" or name.startswith("prior-"):
            prior.append(Path(where).expanduser())
        else:
            folders.setdefault(name, Path(where).expanduser())
    for required in ("rules", "close", "package"):
        if required not in folders:
            raise Bad(f"{path}: the Context has no '{required}' source")
    rules_file = folders["rules"] / RULES_MD
    if not rules_file.is_file():
        raise Bad(f"no {RULES_MD} in the rules folder {folders['rules']}")
    rules = parse_rules(rules_file.read_text(encoding="utf-8"))
    for required in ("close", "package"):
        if not folders[required].is_dir():
            raise Bad(f"the {required} folder {folders[required]} does not exist")
    return Company(str(data.get("name") or path.stem), rules_file, rules, folders["close"],
                   folders["package"], prior)


# --------------------------------------------------------------------------- the period

@dataclass
class Period:
    company: Company
    period: str
    period_dir: Path

    @property
    def rules(self) -> Rules:
        return self.company.rules

    @property
    def work(self) -> Path:
        return self.period_dir / "work"

    @property
    def close_month(self) -> Path:
        return self.company.close / self.period[:4] / self.period

    def tb_file(self, period=None) -> Path:
        p = period or self.period
        return self.company.close / p[:4] / p / "work" / "source" / f"trial-balance-{p}.json"

    def named(self, pattern):
        return self.period_dir / fill(pattern, self.period) if pattern else None

    @property
    def deck(self) -> Path:
        return self.named(self.rules.deck_file)

    @property
    def review_note(self) -> Path:
        return self.named(self.rules.review_note)

    @property
    def tieout_ledger(self) -> Path:
        return self.period_dir / f"tieout-ledger-{self.period}.csv"

    @property
    def tieout_json(self) -> Path:
        return self.work / f"tieout-{self.period}.json"

    @property
    def evidence_file(self) -> Path:
        return self.work / f"BOARD-EVIDENCE-{self.period}.csv"

    def package_files(self) -> dict:
        """The files that make the package, by role; None where the rules keep none."""
        r = self.rules
        return {"deck": self.deck, "script": self.named(r.script_file),
                "cover email": self.named(r.cover_email), "lender pack": self.named(r.lender_pack)}

    def figure_files(self) -> list:
        return [p for p in self.package_files().values() if p is not None]

    def template(self):
        if not self.rules.deck_template:
            return None
        path = Path(fill(self.rules.deck_template, self.period)).expanduser()
        return path if path.is_absolute() else self.company.package / path

    def results_candidates(self) -> list:
        """Files the Results file glob matches, highest v<n> first, then newest."""
        if not self.close_month.is_dir():
            return []
        found = [p for p in self.close_month.glob(fill(self.rules.results_file, self.period)) if p.is_file()]

        def key(p):
            versions = [int(v) for v in VERSION_RE.findall(p.stem)]
            return (max(versions) if versions else 0, p.stat().st_mtime)
        return sorted(found, key=key, reverse=True)

    def results_file(self):
        found = self.results_candidates()
        return found[0] if found else None

    def supporting(self) -> list:
        out = []
        if self.close_month.is_dir():
            for pattern in self.rules.supporting_files:
                for p in sorted(self.close_month.glob(fill(pattern, self.period))):
                    if p.is_file() and p not in out:
                        out.append(p)
        return out

    def _prior_globs(self, period) -> list:
        out = []
        roots = [self.company.package] + self.company.prior
        for pattern in self.rules.previous_packages:
            filled = fill(pattern, period) if period else re.sub(r"\*{2,}", "*", re.sub(r"\{[^}]+\}", "*", pattern))
            if Path(filled).expanduser().is_absolute():
                found = [Path(p) for p in sorted(glob.glob(str(Path(filled).expanduser())))]
            else:
                try:
                    found = [p for root in roots if root.is_dir() for p in sorted(root.glob(filled))]
                except ValueError as exc:
                    raise Bad(f"{RULES_MD}: Previous packages {pattern!r} is not a usable glob: {exc}") from None
            out += [p for p in found if p.is_file() and p not in out]
        return out

    def people_package(self):
        """A package a person made for this period: matched by Previous packages, not ours."""
        ours = {p.resolve() for p in self.figure_files() if p.exists()}
        for p in self._prior_globs(self.period):
            if self.period_dir in p.parents and p.name.startswith("tieout-ledger-"):
                continue
            if p.resolve() not in ours and self.work not in p.parents:
                return p
        return None

    def previous_packages(self, limit=3) -> list:
        """The newest earlier packages, for the format."""
        found = [p for p in self._prior_globs(None) if self.period not in p.name]
        return sorted(found, key=lambda p: p.stat().st_mtime, reverse=True)[:limit]


def resolve(args) -> Period:
    """The company and period an argparse namespace names (company, period, as_of,
    period_dir, contexts_dir)."""
    company = load_company(args.company, args.contexts_dir)
    period = resolve_period(args.period, args.as_of)
    folder = (Path(blank(args.period_dir)).expanduser() if blank(args.period_dir)
              else company.package / fill(company.rules.period_folder, period))
    return Period(company, period, folder)


# --------------------------------------------------------------------------- CSV files and the evidence file

def read_csv(path: Path) -> list:
    """A CSV's rows with stripped values; a missing file is no rows; a BOM is dropped."""
    if not path.is_file():
        return []
    text = path.read_text(encoding="utf-8").removeprefix("\ufeff")
    return [{k: (v or "").strip() for k, v in row.items() if k} for row in csv.DictReader(io.StringIO(text, newline=""))]


def to_csv(rows, columns) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(columns), extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({c: row.get(c, "") for c in columns})
    return buffer.getvalue()


def load_json(path: Path, what: str):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise Bad(f"no {what} at {path}") from None
    except (OSError, ValueError) as exc:
        raise Bad(f"{what} at {path} could not be read: {exc}") from None


def evidence_rows(path: Path) -> dict:
    """The evidence file's rows by id; the first row of an id wins."""
    out = {}
    for row in read_csv(path):
        if row.get("id"):
            out.setdefault(row["id"], row)
    return out


def upsert_row(path: Path, columns, key: str, values: dict) -> str:
    """Create or update the row with id `key`. Other rows keep their exact text, the header
    must be the declared one, no row is ever deleted, and the file is replaced atomically.
    Returns 'created' or 'updated'."""
    raw = path.read_bytes() if path.is_file() else b""
    bom = "\ufeff" if raw.startswith(b"\xef\xbb\xbf") else ""
    text = raw.decode("utf-8").removeprefix("\ufeff")
    lines = text.splitlines(keepends=True)
    records, used = [], 0                      # (cells, the exact text they came from)
    reader = csv.reader(iter(lines))
    for cells in reader:
        records.append((cells, "".join(lines[used:reader.line_num])))
        used = reader.line_num
    end = "\r\n" if records and records[0][1].endswith("\r\n") else "\n"

    def line(cells):
        buffer = io.StringIO()
        csv.writer(buffer, lineterminator=end).writerow(cells)
        return buffer.getvalue()

    if records and [c.strip() for c in records[0][0]] != list(columns):
        raise Refused(f"{path.name} does not start with the header {','.join(columns)}")
    if not records:
        records = [(list(columns), line(columns))]
    index = next((i for i, (cells, _) in enumerate(records) if i and cells and cells[0].strip() == key), None)
    old = dict(zip(columns, [c.strip() for c in records[index][0]])) if index else {}
    row = {c: old.get(c, "") for c in columns}
    row.update({k: "" if v is None else str(v) for k, v in values.items()})
    row["id"] = key
    new = line([row[c] for c in columns])
    if index:
        records[index] = (records[index][0], new)
    else:
        if not records[-1][1].endswith(("\n", "\r")):
            records[-1] = (records[-1][0], records[-1][1] + end)
        records.append(([], new))
    write_atomic(path, bom + "".join(t for _, t in records))
    return "updated" if index else "created"


def write_atomic(path: Path, text: str) -> None:
    """Write through a synced temporary file and a rename, keeping the old file's permissions."""
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile("w", dir=path.parent, prefix=f".{path.name}.", suffix=".tmp",
                                         delete=False, encoding="utf-8", newline="")
    try:
        with handle:
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        if path.exists():
            os.chmod(handle.name, path.stat().st_mode & 0o777)
        os.replace(handle.name, path)
    except BaseException:
        Path(handle.name).unlink(missing_ok=True)
        raise


# --------------------------------------------------------------------------- figures in a file

UNIT_HINTS = [(re.compile(r"\$\s?k\b|\(\s?\$?000s?\s?\)|\$000|\bthousands\b", re.I), "k"),
              (re.compile(r"\$\s?mm?\b|\bmillions\b", re.I), "M"),
              (re.compile(r"\(\s?\$\s?\)"), "1")]


def unit_hint(texts):
    """The unit a table's header or caption names, or None."""
    for text in texts:
        for rx, unit in UNIT_HINTS:
            if rx.search(text or ""):
                return unit
    return None


def figure(file, location, where, p: Printed, unit, context="") -> dict:
    shown = "%" if p.percent else ({"k": "k", "m": "M", "b": "B"}[p.suffix] if p.suffix else unit or "")
    return {"file": file, "location": location, "where": where, "printed": p.text, "value": p.value,
            "decimals": p.decimals, "suffix": p.suffix, "percent": p.percent, "unit": shown, "context": context}


def from_cell(file, location, where, text, unit, context, ignore) -> list:
    """A cell holding one number is a figure (a bare year excepted); otherwise its marked numbers."""
    single = parse_printed((text or "").strip())
    if single is not None:
        if any(rx.search(single.text) for rx in ignore) or is_year(single):
            return []
        return [figure(file, location, where, single, unit, context)]
    return [figure(file, location, where, p, unit, context) for p in tokens(text, ignore) if figure_in_text(p)]


def from_text(file, location, where, text, ignore) -> list:
    return [figure(file, location, where, p, None) for p in tokens(text, ignore) if figure_in_text(p)]


class _Tables(HTMLParser):
    """An HTML file's tables (rows of (text, is_header) with a caption) and its visible text outside them."""

    SKIP = {"script", "style", "template", "head", "title"}
    BLOCK = {"br", "p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "section", "header", "footer"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.text, self._stack = [], [], []
        self._cell, self._header, self._caption, self._skip = None, False, None, 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "table":
            self._stack.append({"rows": [], "caption": "", "row": None})
        elif tag == "caption" and self._stack:
            self._caption = []
        elif tag == "tr" and self._stack:
            self._stack[-1]["row"] = []
        elif tag in ("td", "th") and self._stack:
            self._cell, self._header = [], tag == "th"
        elif tag in self.BLOCK and not self._stack:
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in ("td", "th") and self._stack and self._cell is not None:
            if self._stack[-1]["row"] is not None:
                self._stack[-1]["row"].append((" ".join("".join(self._cell).split()), self._header))
            self._cell = None
        elif tag == "tr" and self._stack and self._stack[-1]["row"] is not None:
            if self._stack[-1]["row"]:
                self._stack[-1]["rows"].append(self._stack[-1]["row"])
            self._stack[-1]["row"] = None
        elif tag == "caption" and self._caption is not None and self._stack:
            self._stack[-1]["caption"] = " ".join("".join(self._caption).split())
            self._caption = None
        elif tag == "table" and self._stack:
            table = self._stack.pop()
            self.tables.append({"rows": table["rows"], "caption": table["caption"]})
        elif tag in self.BLOCK and not self._stack:
            self.text.append("\n")

    def handle_data(self, data):
        if self._skip:
            return
        if self._caption is not None:
            self._caption.append(data)
        elif self._cell is not None:
            self._cell.append(data)
        elif not self._stack:
            self.text.append(data)


def read_html(path: Path) -> _Tables:
    parser = _Tables()
    parser.feed(re.sub(r"<!--.*?-->", " ", path.read_text(encoding="utf-8", errors="replace"), flags=re.S))
    parser.close()
    return parser


def html_visible_text(path: Path) -> str:
    """Everything a reader sees, tables included, one line per block or row."""
    parsed = read_html(path)
    lines = ["".join(parsed.text)]
    for table in parsed.tables:
        lines.append(table["caption"])
        lines += [" | ".join(text for text, _ in row) for row in table["rows"]]
    return "\n".join(lines)


def table_header(rows):
    """The header row's texts and the index of the first data row."""
    if not rows:
        return None, 0
    if all(is_th for _, is_th in rows[0]):
        return [t for t, _ in rows[0]], 1
    if len(rows[0]) > 1 and all(parse_printed(t) is None for t, _ in rows[0][1:]):
        return [t for t, _ in rows[0]], 1
    return None, 0


def html_figures(path: Path, unit, ignore=()) -> list:
    parsed = read_html(path)
    out = []
    for t_no, table in enumerate(parsed.tables, start=1):
        header, first = table_header(table["rows"])
        t_unit = unit_hint([table["caption"]] + (header or [])) or unit
        for r_no, row in enumerate(table["rows"][first:], start=first + 1):
            label = row[0][0] if row else ""
            for c_no, (text, _) in enumerate(row):
                if c_no == 0 and parse_printed(text) is None:
                    continue
                col = header[c_no] if header and c_no < len(header) else f"col {c_no + 1}"
                out += from_cell(path.name, f"table {t_no} row {r_no} col {c_no + 1}", "table", text, t_unit,
                                 f"{label} | {col}" if c_no else label, ignore)
    for line in "".join(parsed.text).splitlines():
        out += from_text(path.name, "text", "text", line, ignore)
    return out


def results_rows(path: Path, period: str, unit, ignore=()) -> list:
    """The results file's table figures as results-figures rows: metric '<row> | <column>'
    (a repeat in a later table gets ' (table N)'), value in base units."""
    rows, first_table = [], {}
    for f in html_figures(path, unit, ignore):
        if f["where"] != "table" or " | " not in f["context"]:
            continue
        p = parse_printed(f["printed"])
        if p is None:
            continue
        table_no = int(f["location"].split()[1])
        key = f["context"].lower()
        first_table.setdefault(key, table_no)
        metric = f["context"] if first_table[key] == table_no else f"{f['context']} (table {table_no})"
        value = dollars(p, f["unit"] if f["unit"] != "%" else None)
        rows.append({"id": f"R{len(rows) + 1}", "file": path.name, "period": period, "metric": metric,
                     "value": repr(round(value, 6)), "unit": f["unit"] or "1", "printed": f["printed"],
                     "location": f["location"]})
    return rows


def _plain(value: float) -> str:
    return f"{value:.6f}".rstrip("0").rstrip(".")


def _shapes(shapes) -> list:
    out = []
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6:  # a group: look inside it
            out += _shapes(shape.shapes)
        else:
            out.append(shape)
    return out


def open_deck(path: Path):
    from pptx import Presentation
    try:
        return Presentation(str(path))
    except Exception as exc:  # any deck python-pptx cannot read
        raise Bad(f"{path.name} could not be read as a deck: {type(exc).__name__}: {exc}") from None


def pptx_figures(path: Path, unit, ignore=()) -> list:
    out = []
    for s_no, slide in enumerate(open_deck(path).slides, start=1):
        loc = f"slide {s_no}"
        shapes = _shapes(slide.shapes)
        slide_texts = [sh.text_frame.text for sh in shapes if getattr(sh, "has_text_frame", False)]
        for shape in shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                header = rows[0] if rows else []
                t_unit = unit_hint(header + slide_texts) or unit
                for r_no, row in enumerate(rows):
                    for c_no, text in enumerate(row):
                        if c_no == 0 and parse_printed(text) is None:
                            continue
                        col = header[c_no] if r_no and c_no < len(header) else ""
                        context = f"{row[0]} | {col}" if (c_no and col) else row[0]
                        out += from_cell(path.name, loc, "table", text, t_unit, context, ignore)
            elif getattr(shape, "has_chart", False):
                t_unit = unit_hint(slide_texts) or unit
                for plot in shape.chart.plots:
                    for series in plot.series:
                        for value in series.values:
                            p = parse_printed(_plain(float(value))) if value is not None else None
                            if p is not None:
                                out.append(figure(path.name, loc, "chart", p, t_unit, series.name or ""))
            elif getattr(shape, "has_text_frame", False):
                for line in shape.text_frame.text.splitlines():
                    out += from_text(path.name, loc, "text", line, ignore)
        if slide.has_notes_slide:
            for line in slide.notes_slide.notes_text_frame.text.splitlines():
                out += from_text(path.name, loc, "notes", line, ignore)
    return out


def pptx_text(path: Path) -> str:
    parts = []
    for slide in open_deck(path).slides:
        for shape in _shapes(slide.shapes):
            if getattr(shape, "has_table", False):
                parts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
            elif getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def xlsx_figures(path: Path, unit, ignore=()) -> list:
    """Every numeric cell, printed at the decimals its number format shows."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as exc:  # any workbook openpyxl cannot read
        raise Bad(f"{path.name} could not be read as a workbook: {type(exc).__name__}: {exc}") from None
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value, where = cell.value, f"{ws.title}!{cell.coordinate}"
                if isinstance(value, bool) or value is None:
                    continue
                if isinstance(value, (int, float)):
                    fmt = cell.number_format or ""
                    m = re.search(r"0\.(0+)", fmt)
                    decimals = len(m.group(1)) if m else (0 if fmt != "General" and "0" in fmt else None)
                    if "%" in fmt:
                        printed = f"{float(value) * 100:.{decimals or 0}f}%"
                    elif decimals is not None:
                        printed = f"{float(value):.{decimals}f}"
                    else:
                        printed = _plain(float(value))
                    p = parse_printed(printed)
                    if p is not None and not is_year(p) and not any(rx.search(p.text) for rx in ignore):
                        out.append(figure(path.name, where, "cell", p, unit))
                elif isinstance(value, str):
                    out += from_cell(path.name, where, "cell", value, unit, "", ignore)
    return out


MD_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
MD_RULE_CELL = re.compile(r"^:?-{2,}:?$")


def text_figures(path: Path, unit, ignore=()) -> list:
    """Running text and Markdown tables, by line; HTML comments are not read."""
    out, header = [], None
    text = path.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"<!--.*?-->", lambda m: "\n" * m.group(0).count("\n"), text, flags=re.S)
    for n, line in enumerate(text.splitlines(), start=1):
        if not MD_TABLE_ROW.match(line):
            header = None
            out += from_text(path.name, f"line {n}", "text", line, ignore)
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and all(MD_RULE_CELL.match(c.replace(" ", "")) for c in cells if c):
            continue
        if header is None:
            header = cells
            for c in cells:
                p = parse_printed(c)
                if p is not None and not is_year(p):
                    out.append(figure(path.name, f"line {n}", "table", p, unit_hint(header) or unit))
            continue
        t_unit = unit_hint(header) or unit
        for c_no, c in enumerate(cells):
            if c_no == 0 and parse_printed(c) is None:
                continue
            col = header[c_no] if c_no < len(header) else ""
            out += from_cell(path.name, f"line {n}", "table", c, t_unit,
                             f"{cells[0]} | {col}" if c_no else cells[0], ignore)
    return out


def file_text(path: Path) -> str:
    """The visible text of a package file, for the placeholder and prose tests."""
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return pptx_text(path)
    if suffix in (".html", ".htm"):
        return html_visible_text(path)
    if suffix == ".xlsx":
        return ""
    return re.sub(r"<!--.*?-->", " ", path.read_text(encoding="utf-8", errors="replace"), flags=re.S)


def extract(path: Path, unit=None, ignore=()) -> list:
    """Every figure printed in a .pptx, .html, .xlsx, .md or .txt file, with where it is."""
    if not path.is_file():
        raise Bad(f"no file {path}")
    suffix = path.suffix.lower()
    readers = {".pptx": pptx_figures, ".html": html_figures, ".htm": html_figures, ".xlsx": xlsx_figures,
               ".md": text_figures, ".txt": text_figures}
    if suffix not in readers:
        raise Bad(f"{path.name}: cannot read figures from a {suffix or 'suffixless'} file "
                  f"(.pptx, .html, .xlsx, .md, .txt)")
    return readers[suffix](path, unit, ignore)


def match_location(name: str, location: str) -> str:
    """The part of a location a ledger row must match: the slide of a deck, the cell of a
    workbook, nothing for text and HTML (matched anywhere in the file)."""
    suffix, loc = Path(name).suffix.lower(), (location or "").strip()
    if suffix == ".pptx":
        m = re.match(r"(?i)slide\s*(\d+)", loc)
        return f"slide {m.group(1)}" if m else loc.lower()
    if suffix == ".xlsx":
        return loc.replace("$", "").lower()
    return ""


# --------------------------------------------------------------------------- the close

def close_state(per: Period) -> dict:
    """Whether the close is done by the rules' Close gate. The month-end-check gate runs the
    month-end-workstream skill's check script on the close folder."""
    rules = per.rules
    attested = per.period in rules.closed_by_hand
    state = {"gate": rules.close_gate, "attested": attested, "done": attested, "check": None}
    if rules.close_gate == "month-end-check":
        script = Path(MONTH_END_CHECK).expanduser()
        try:
            if not script.is_file():
                raise FileNotFoundError(f"month-end-check is not installed at {script}")
            done = subprocess.run([sys.executable, str(script), str(per.company.close), "--period", per.period,
                                   "--format", "json"], capture_output=True, text=True, timeout=300)
            if done.returncode != 0:
                raise RuntimeError((done.stderr or done.stdout).strip().splitlines()[-1:] or "failed")
            result = json.loads(done.stdout)
            state["check"] = {"met": result["met"], "of": len(result["tests"]), "done": result["done"],
                              "open": [k for k, t in result["tests"].items() if not t["met"]]}
            state["done"] = attested or bool(result["done"])
        except Exception as exc:  # a missing month folder or check is a state, not a crash
            state["check"] = {"error": f"{type(exc).__name__}: {exc}"}
    check = state["check"] or {}
    if state["done"]:
        state["why"] = ("closed by hand (Closed by hand in the rules)" if attested and not check.get("done")
                        else "month-end-check: every test met")
    elif rules.close_gate == "attested":
        state["why"] = f"{per.period} is not in the rules' Closed by hand list"
    else:
        state["why"] = check.get("error") or (f"month-end-check: {check.get('met', 0)} of {check.get('of', 4)} "
                                              f"tests met (open: {', '.join(check.get('open', []))})")
    return state


# --------------------------------------------------------------------------- the tie-out

SOURCE_RE = re.compile(r"^(?P<neg>-)?\s*(?P<kind>results|tb-month|tb-ytd|tb|calc|text|person)\s*:(?P<arg>.*)$", re.I)
TIEOUT_COLUMNS = ("file", "period", "metric", "value", "location", "id", "source", "exact", "state")
OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv,
       ast.USub: operator.neg, ast.UAdd: operator.pos}


class SourceError(Exception):
    """A source that cannot be computed: the figure fails."""


def parse_source(text):
    """(negated, kind, argument) of a ledger source such as '-tb-month:40000-49999'."""
    m = SOURCE_RE.match((text or "").strip())
    if not m:
        raise SourceError(f"source {text!r} is not results:, tb:, tb-month:, tb-ytd:, calc:, text: or person:")
    return bool(m.group("neg")), m.group("kind").lower(), m.group("arg").strip()


def account_spec(text: str) -> list:
    """'10400', '15000-15999' or '19300, 19350' as (low, high) pairs."""
    out = []
    for part in re.split(r"[,;]", text.replace("`", "")):
        part = part.strip()
        m = re.fullmatch(r"(\d+)\s*[-\u2013]\s*(\d+)", part)
        if m and int(m.group(1)) <= int(m.group(2)):
            out.append((m.group(1), m.group(2)))
        elif part:
            out.append((part, part))
    return out


def in_spec(account: str, spec) -> bool:
    return any(account == low if low == high else (account.isdigit() and int(low) <= int(account) <= int(high))
               for low, high in spec)


def evaluate(expr: str, value_of) -> float:
    """calc: arithmetic over numbers and figure ids, nothing else."""
    try:
        tree = ast.parse(expr, mode="eval")
    except SyntaxError:
        raise SourceError(f"calc:{expr}: not an arithmetic expression") from None

    def walk(node):
        if isinstance(node, ast.Expression):
            return walk(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
            return float(node.value)
        if isinstance(node, ast.Name) and re.fullmatch(r"F\d+", node.id):
            return value_of(node.id)
        if isinstance(node, ast.BinOp) and type(node.op) in OPS:
            left, right = walk(node.left), walk(node.right)
            if isinstance(node.op, ast.Div) and right == 0:
                raise SourceError(f"calc:{expr}: division by zero")
            return OPS[type(node.op)](left, right)
        if isinstance(node, ast.UnaryOp) and type(node.op) in OPS:
            return OPS[type(node.op)](walk(node.operand))
        raise SourceError(f"calc:{expr}: only numbers, figure ids, + - * / and parentheses")
    return walk(tree)


class Books:
    """What a figure is computed from: the results figures, the results text and the trial balance pulls."""

    def __init__(self, per: Period):
        self.per = per
        self.results = read_csv(per.work / RESULTS_FIGURES_CSV)
        self._tb, self._text = {}, None

    def tb(self, p: str) -> dict:
        if p not in self._tb:
            try:
                self._tb[p] = json.loads(self.per.tb_file(p).read_text(encoding="utf-8"))
            except (OSError, ValueError):
                self._tb[p] = None
        if not self._tb[p]:
            raise SourceError(f"no trial balance pull for {p} ({self.per.tb_file(p)})")
        return self._tb[p]

    def tb_sum(self, kind: str, spec_text: str) -> float:
        spec = account_spec(spec_text)
        if not spec:
            raise SourceError(f"{kind}: names no account")

        def total(p, column):
            hit = [r for r in self.tb(p).get("rows") or [] if in_spec(str(r.get("Account", "")).strip(), spec)]
            if not hit:
                raise SourceError(f"{kind}:{spec_text}: no account in the {p} trial balance")
            return sum(float(r.get(column) or 0) for r in hit)

        period = self.per.period
        if kind == "tb":
            return total(period, "Ending balance")
        this = total(period, "YTD activity")
        if kind == "tb-ytd":
            return this
        # tb-month: this period's year-to-date less the prior period's, within one fiscal year.
        start = self.tb(period).get("fiscal_year_start") or f"{period[:4]}-01-01"
        if str(start)[:7] == period:
            return this
        try:
            prior_tb = self.tb(prior_period(period))
        except SourceError:
            raise SourceError(f"tb-month needs the {prior_period(period)} trial balance pull too") from None
        if (prior_tb.get("fiscal_year_start") or start) != start:
            return this
        try:
            return this - total(prior_period(period), "YTD activity")
        except SourceError:
            return this

    def result(self, arg: str) -> float:
        if not self.results:
            raise SourceError(f"no {RESULTS_FIGURES_CSV}: run board_package_pack.py")
        key = arg.strip()
        if re.fullmatch(r"R\d+", key):
            rows = [r for r in self.results if r.get("id") == key]
        else:
            rows = [r for r in self.results if r.get("metric", "").strip().lower() == key.lower()]
        if not rows:
            raise SourceError(f"results:{key}: no such figure in {RESULTS_FIGURES_CSV}")
        return float(rows[0]["value"])

    def results_text(self) -> list:
        if self._text is None:
            path = self.per.results_file()
            if path is None:
                text = ""
            elif path.suffix.lower() in (".html", ".htm"):
                text = html_visible_text(path)
            else:
                text = path.read_text(encoding="utf-8", errors="replace")
            self._text = [p for p in tokens(text) if figure_in_text(p)]
        return self._text


def _num(value: float) -> str:
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text


def tieout(per: Period) -> dict:
    """Every figure-ledger row computed from its source and judged: ties, differs, error or
    stated (by a person). An approved exception lets a failing figure pass."""
    rules = per.rules
    rows = read_csv(per.work / PACKAGE_FIGURES_CSV)
    if not rows:
        raise Bad(f"no figure ledger at {per.work / PACKAGE_FIGURES_CSV}")
    missing = [c for c in ("id", "file", "printed", "source") if c not in rows[0]]
    if missing:
        raise Bad(f"{PACKAGE_FIGURES_CSV} needs the columns {', '.join(PACKAGE_COLUMNS)}; missing {', '.join(missing)}")
    books = Books(per)
    by_id = {r["id"]: r for r in rows if r.get("id")}
    exceptions = {r["item"]: r.get("note", "") for r in evidence_rows(per.evidence_file).values()
                  if r.get("kind") == "exception" and r.get("state") == "approved"}
    computed, errors, visiting = {}, {}, []

    def value_of(fid):
        if computed.get(fid) is not None:
            return computed[fid]
        if fid in errors:
            raise SourceError(f"{fid} cannot be computed ({errors[fid]})")
        if fid not in by_id:
            raise SourceError(f"no figure {fid} in the ledger")
        if fid in visiting:
            raise SourceError(f"calc refers to itself through {' -> '.join(visiting + [fid])}")
        visiting.append(fid)
        try:
            computed[fid] = source_value(by_id[fid])
        finally:
            visiting.pop()
        return computed[fid]

    def source_value(row):
        neg, kind, arg = parse_source(row.get("source", ""))
        p = parse_printed(row.get("printed", ""))
        if p is None:
            raise SourceError(f"printed {row.get('printed')!r} is not a number")
        if kind == "results":
            value = books.result(arg)
        elif kind.startswith("tb"):
            value = books.tb_sum(kind, arg)
        elif kind == "calc":
            value = evaluate(arg, value_of)
        elif kind == "text":
            # The results file's prose prints an amount that rounds to this figure.
            target = dollars(p, row.get("unit"))
            for t in books.results_text():
                candidate = (-1 if neg else 1) * dollars(t, None if t.suffix or t.percent else rules.results_unit)
                if abs(round_to_print(candidate, p, row.get("unit")) - target) <= rules.tolerance:
                    return candidate
            raise SourceError("text: the results file prints no such amount")
        else:
            raise SourceError("person: stated, not computed")
        return -value if neg else value

    figures = []
    for row in rows:
        fid = row.get("id", "")
        out = {k: row.get(k, "") for k in ("id", "file", "location", "printed", "unit", "metric", "source")}
        p = parse_printed(row.get("printed", ""))
        try:
            if not blank(row.get("source")):
                raise SourceError("no source")
            if p is None:
                raise SourceError(f"printed {row.get('printed')!r} is not a number")
            _, kind, _ = parse_source(row["source"])
            out["printed_value"] = dollars(p, row.get("unit"))
            if kind == "person":
                out.update(state="stated", detail="stated by a person; name it in the review note")
            else:
                exact = value_of(fid)
                rounded = round_to_print(exact, p, row.get("unit"))
                gap = out["printed_value"] - rounded
                out.update(exact=exact, rounded=rounded, gap=round(gap, 6))
                out["state"] = "ties" if abs(gap) <= rules.tolerance + 1e-6 else "differs"
                if out["state"] == "differs":
                    out["detail"] = (f"printed {row.get('printed')} ({out['printed_value']:,.2f}); books {exact:,.2f}, "
                                     f"{rounded:,.2f} at the printed precision")
        except (SourceError, Bad) as exc:
            errors.setdefault(fid, str(exc))
            out.update(state="error", detail=str(exc))
        if out["state"] in ("differs", "error") and fid in exceptions:
            out["exception"] = exceptions[fid]
        figures.append(out)

    ledger = []
    for f in figures:
        if "rounded" not in f:
            continue
        common = {"period": per.period, "metric": f"{f['id']} {f['metric']}".strip(), "id": f["id"],
                  "source": f["source"], "state": f["state"]}
        ledger.append({**common, "file": f["file"] or "(not printed)", "value": _num(f["printed_value"]),
                       "location": f["location"], "exact": ""})
        ledger.append({**common, "file": "books", "value": _num(f["rounded"]), "location": f["source"],
                       "exact": _num(f["exact"])})
    # The same comparison report-tieout makes on this ledger: package against books, per metric.
    findings = []
    for i in range(0, len(ledger), 2):
        spread = round(abs(float(ledger[i]["value"]) - float(ledger[i + 1]["value"])), 2)
        if spread > rules.tolerance:
            # Each side with its value and where it came from: the slide or cell, and the source.
            sides = sorted(ledger[i:i + 2], key=lambda r: r["file"])
            detail = ", ".join(f"{r['file']} {float(r['value']):,.0f}" + (f" ({r['location']})" if r["location"] else "")
                               for r in sides)
            findings.append({"check": "disagree", "metric": ledger[i]["metric"], "period": per.period,
                             "spread": spread, "detail": detail})
    kit_findings = [k for k in findings if k["metric"].split(" ", 1)[0] not in exceptions]
    failing = [f["id"] for f in figures if f["state"] in ("differs", "error") and f["id"] not in exceptions]
    counts = {}
    for f in figures:
        counts[f["state"]] = counts.get(f["state"], 0) + 1
    return {"company": per.company.name, "period": per.period, "figures": figures, "counts": counts,
            "exceptions": sorted(exceptions), "failing": failing,
            "report_tieout": {"findings": findings, "passed": not kit_findings},
            "passed": not failing and not kit_findings, "ledger_rows": ledger}


def tieout_ledger_text(result: dict) -> str:
    return to_csv(result["ledger_rows"], TIEOUT_COLUMNS)
