# /// script
# dependencies = ["openpyxl"]
# ///
"""Read a per-finished-good P&L workbook into two long tables and tie revenue to the TB.

The workbook has a tab with one row per product and one block of columns per month (and
optionally a year-to-date tab of the same shape), and a trial-balance tab with categories
down the side and months across. Where each of those sits is described by a layout file
(see _common.py), given with --layout or the `cost_model_layout` setting in
[product-costing-workstream]; --month-tab, --ytd-tab and --tb-tab override its tab names.

Outputs:
  fct_fg_pl   one row per product x period x basis (month or ytd)
  fct_tb      one row per trial-balance line x month
  reconciliation  product revenue (month basis) against TB revenue, within max($1, 0.05%)

A product with revenue and zero standard cost is flagged cost_pending and its margin
percentages are null. Without --out the whole result is printed as JSON; with --out the
two tables are written as CSV with reconciliation.json beside them, and the exit code is 1
when the revenue tie is outside tolerance. Read-only.

Example:
  python3 fg_pl_extract.py "Cost model rev3.xlsx" --year 2026 --layout layout.json --out work/
"""

import argparse
import csv
import json
import sys
from pathlib import Path

from _common import fail, find_tab, load_layout, parse_product_tab, parse_tb


def reconcile(fg, tb, revenue_category):
    """Sum of product revenue (month basis) against the TB revenue category."""
    fg_rev = sum((r.get("sales_amount") or 0.0) for r in fg if r["basis"] == "month")
    tb_rev = sum(r["amount"] for r in tb if (r.get("category") or "").strip().lower() == revenue_category.strip().lower())
    tol = max(1.0, 0.0005 * abs(fg_rev))
    return {"fg_revenue": round(fg_rev, 2), "tb_revenue": round(tb_rev, 2), "difference": round(fg_rev - tb_rev, 2),
            "tolerance": round(tol, 2), "status": "green" if abs(fg_rev - tb_rev) <= tol else "red"}


def extract(workbook, year, layout):
    import openpyxl

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    tb_name = find_tab(wb.sheetnames, layout["tb_tab"])
    missing = [t for t, found in ((layout["month_tab"], layout["month_tab"] in wb.sheetnames), (layout["tb_tab"], tb_name)) if not found]
    if missing:
        raise ValueError(f"Tabs not found: {missing}. Sheets: {wb.sheetnames}")
    fg = parse_product_tab(wb[layout["month_tab"]], "month", year, layout)
    ytd = layout.get("ytd_tab")
    if ytd and ytd in wb.sheetnames:
        fg += parse_product_tab(wb[ytd], "ytd", year, layout)
    else:
        ytd = None
    tb = parse_tb(wb[tb_name], year, layout)
    return {"workbook": str(workbook), "year": year, "tabs": {"month": layout["month_tab"], "ytd": ytd, "tb": tb_name},
            "fct_fg_pl": fg, "fct_tb": tb, "products": len({r["product_fg"] for r in fg}),
            "cost_pending": sum(1 for r in fg if r.get("cost_pending")),
            "reconciliation": reconcile(fg, tb, layout["revenue_category"])}


def write_tables(result, out_dir):
    out = Path(out_dir).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    for name in ("fct_fg_pl", "fct_tb"):
        rows = result[name]
        with (out / f"{name}.csv").open("w", newline="", encoding="utf-8") as fh:
            if rows:
                fields = list(dict.fromkeys(k for r in rows for k in r))
                w = csv.DictWriter(fh, fieldnames=fields)
                w.writeheader()
                w.writerows(rows)
    (out / "reconciliation.json").write_text(json.dumps(result["reconciliation"], indent=1), encoding="utf-8")
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="Extract fact tables from a per-finished-good P&L workbook and tie product revenue to the TB.")
    ap.add_argument("workbook")
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--layout", help="Layout JSON (default: the cost_model_layout setting)")
    ap.add_argument("--month-tab", help="Override the layout's month tab")
    ap.add_argument("--ytd-tab", help="Override the layout's year-to-date tab")
    ap.add_argument("--tb-tab", help="Override the layout's trial-balance tab (prefix match)")
    ap.add_argument("--revenue-category", help="Override the layout's TB revenue category")
    ap.add_argument("--out", help="Write fct_fg_pl.csv, fct_tb.csv and reconciliation.json here instead of JSON to stdout")
    a = ap.parse_args(argv)

    path = Path(a.workbook).expanduser()
    if not path.is_file():
        fail(f"Not a file: {path}")
    layout = load_layout(a.layout, {"month_tab": a.month_tab, "ytd_tab": a.ytd_tab, "tb_tab": a.tb_tab, "revenue_category": a.revenue_category})
    try:
        result = extract(path, a.year, layout)
    except ValueError as exc:
        fail(str(exc))
    if not a.out:
        print(json.dumps(result, indent=1, default=str))
        return 0
    out = write_tables(result, a.out)
    rec = result["reconciliation"]
    print(f"{len(result['fct_fg_pl'])} FG rows ({result['products']} products, {result['cost_pending']} cost-pending), {len(result['fct_tb'])} TB rows -> {out}")
    print(f"tie-out {rec['status'].upper()}: FG revenue {rec['fg_revenue']:,.0f} vs TB {rec['tb_revenue']:,.0f} (diff {rec['difference']:,.0f}, tol {rec['tolerance']:,.0f})")
    return 0 if rec["status"] == "green" else 1


if __name__ == "__main__":
    sys.exit(main())
