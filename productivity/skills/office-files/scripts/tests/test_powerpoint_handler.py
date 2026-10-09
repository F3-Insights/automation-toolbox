import json

from pptx import Presentation
from pptx.util import Inches

import powerpoint_handler


def make_deck(path):
    deck = Presentation()
    deck.core_properties.author = "Priya"
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "September results"
    slide.placeholders[1].text = "Revenue up 4% on plan"
    slide.notes_slide.notes_text_frame.text = "Mention the Lakeview Hardware renewal"
    table_slide = deck.slides.add_slide(deck.slide_layouts[5])
    table_slide.shapes.title.text = "Cash"
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Week", "Balance"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "250,000"
    deck.save(path)
    return str(path)


def test_json_reads_titles_text_notes_and_tables(tmp_path, capsys):
    path = make_deck(tmp_path / "board.pptx")
    assert powerpoint_handler.main(["--ppt-file", path]) == 0
    deck = json.loads(capsys.readouterr().out)
    assert deck["slide_count"] == 2 and deck["metadata"]["author"] == "Priya"
    first, second = deck["slides"]
    assert first["title"] == "September results"
    assert "Revenue up 4% on plan" in [s["text"] for s in first["shapes"]]
    assert first["notes"] == "Mention the Lakeview Hardware renewal"
    assert "1 | 250,000" in "\n".join(s["text"] for s in second["shapes"])


def test_markdown_without_notes(tmp_path, capsys):
    path = make_deck(tmp_path / "board.pptx")
    out_file = tmp_path / "board.md"
    powerpoint_handler.main(["--ppt-file", path, "--output-format", "markdown", "--no-notes",
                             "--output", str(out_file)])
    text = capsys.readouterr().out
    assert text.startswith("# board.pptx") and "## Slide 2" in text
    assert "Lakeview" not in text
    assert out_file.read_text() + "\n" == text


def test_notes_page_without_body_box(tmp_path, capsys):
    path = make_deck(tmp_path / "board.pptx")
    deck = Presentation(path)
    notes = deck.slides[0].notes_slide
    body = notes.notes_placeholder
    body._element.getparent().remove(body._element)
    deck.save(path)
    assert powerpoint_handler.main(["--ppt-file", path]) == 0
    assert json.loads(capsys.readouterr().out)["slides"][0]["notes"] is None
