# /// script
# dependencies = ["openpyxl"]
# ///
"""Turn a recurring SAP CO/PP extract into one long table with provenance.

Clients send the same reports each month as workbooks with one tab per month, headers in
English or Spanish, and columns in whatever order the layout variant produced. This
locates columns by header text, never by position, stacks the month tabs, and gives every
output row the file, tab and row it came from and the file's SHA-256.

Extract types and the canonical columns each yields (run with --types to list them):

    cost-breakdown  material, cost_component, amount, currency, plant
    routing         material, group, operation, work_center, control_key, std_value, unit, base_qty
    ppv             material, vendor, po, qty, std_price, actual_price, variance, currency
    scrap           material, order, qty, reason, cost_center, amount
    freight-in      material, vendor, document, amount, currency
    rates           cost_center, activity_type, rate, unit, currency, period
    wc-list         work_center, description, cost_center, plant, capacity_category
    tb              account, description, amount, period, company_code

A client whose headers use other words adds them in a JSON file, {"<type>": {"<column>":
["header", ...]}}, named by --synonyms or the `extract_synonyms` setting in
[product-costing-workstream]; they are added to the built-in ones.

A month tab is any tab whose name holds a month token (Jan, Ene, 2026-03, 03); the token
becomes the `period` column when the extract has none. Tabs without the type's required
headers are listed as skipped, never guessed. When no tab has them the script prints a
DATA REQUEST naming the columns and exits 1.

Reads .xlsx (openpyxl) and .csv. Read-only. Prints JSON, or with --out writes the long
table as CSV and prints a summary.

Example:
  python3 sap_extract_normalize.py ppv PPV_2026.xlsx --out work/ppv_long.csv
"""

import argparse
import csv
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from _common import SKILL, fail, open_csv_text, settings, sha256

# canonical column -> header synonyms (lower-cased, accents stripped, punctuation ignored)
EXTRACTS = {
    "cost-breakdown": {"required": ["material", "cost_component", "amount"], "columns": {
        "material": ["material", "material number", "part", "part number", "articulo", "numero de material", "producto"],
        "cost_component": ["cost component", "cost comp", "component", "componente de costo", "componente", "cost element", "elemento de costo"],
        "amount": ["amount", "value", "total", "importe", "valor", "monto", "costo"],
        "currency": ["currency", "crcy", "moneda"],
        "plant": ["plant", "plnt", "centro"]}},
    "routing": {"required": ["material", "operation", "work_center", "std_value"], "columns": {
        "material": ["material", "material number", "part", "articulo", "producto"],
        "group": ["group", "group counter", "task list group", "grupo", "grupo hoja ruta"],
        "operation": ["operation", "op", "operation number", "operacion", "num operacion"],
        "work_center": ["work center", "work ctr", "wc", "puesto de trabajo", "puesto trabajo", "centro de trabajo"],
        "control_key": ["control key", "ctrl key", "clave de control", "clave control"],
        "std_value": ["std value", "standard value", "labor", "machine", "setup", "valor prefijado", "valor estandar", "minutes", "minutos", "tiempo"],
        "unit": ["unit", "un", "unidad", "unit of measure"],
        "base_qty": ["base quantity", "base qty", "cantidad base", "cant base"]}},
    "ppv": {"required": ["material", "variance"], "columns": {
        "material": ["material", "material number", "part", "articulo"],
        "vendor": ["vendor", "supplier", "proveedor", "acreedor"],
        "po": ["purchase order", "po", "purchasing document", "pedido", "orden de compra"],
        "qty": ["quantity", "qty", "cantidad"],
        "std_price": ["standard price", "std price", "precio estandar", "precio std"],
        "actual_price": ["actual price", "po price", "net price", "precio real", "precio neto"],
        "variance": ["variance", "price variance", "ppv", "diferencia", "variacion", "desviacion de precio"],
        "currency": ["currency", "crcy", "moneda"]}},
    "scrap": {"required": ["material", "qty"], "columns": {
        "material": ["material", "material number", "part", "articulo"],
        "order": ["order", "production order", "orden", "orden de produccion"],
        "qty": ["scrap qty", "scrap quantity", "quantity", "qty", "cantidad", "cantidad rechazo", "desperdicio"],
        "reason": ["reason", "reason code", "motivo", "causa"],
        "cost_center": ["cost center", "cost ctr", "centro de costo", "centro coste"],
        "amount": ["amount", "value", "importe", "valor", "costo"]}},
    "freight-in": {"required": ["amount"], "columns": {
        "material": ["material", "material number", "part", "articulo"],
        "vendor": ["vendor", "carrier", "supplier", "proveedor", "transportista"],
        "document": ["document", "document number", "invoice", "documento", "factura"],
        "amount": ["amount", "freight", "freight amount", "importe", "flete", "valor"],
        "currency": ["currency", "crcy", "moneda"]}},
    "rates": {"required": ["cost_center", "activity_type", "rate"], "columns": {
        "cost_center": ["cost center", "cost ctr", "centro de costo", "centro coste"],
        "activity_type": ["activity type", "acttyp", "activity", "clase de actividad", "tipo de actividad"],
        "rate": ["rate", "price", "fixed price", "variable price", "total price", "tarifa", "precio"],
        "unit": ["unit", "price unit", "unidad"],
        "currency": ["currency", "crcy", "moneda"],
        "period": ["period", "per", "periodo", "mes"]}},
    "wc-list": {"required": ["work_center"], "columns": {
        "work_center": ["work center", "work ctr", "wc", "puesto de trabajo", "centro de trabajo"],
        "description": ["description", "short text", "descripcion", "texto breve"],
        "cost_center": ["cost center", "cost ctr", "centro de costo", "centro coste"],
        "plant": ["plant", "plnt", "centro"],
        "capacity_category": ["capacity category", "cap cat", "categoria de capacidad"]}},
    "tb": {"required": ["account", "amount"], "columns": {
        "account": ["account", "g/l account", "gl account", "cuenta", "cuenta de mayor"],
        "description": ["description", "account name", "short text", "descripcion", "nombre de cuenta"],
        "amount": ["amount", "balance", "period balance", "saldo", "importe", "valor"],
        "period": ["period", "per", "periodo", "mes"],
        "company_code": ["company code", "cocd", "sociedad"]}},
}

MONTHS = {"jan": 1, "ene": 1, "feb": 2, "mar": 3, "apr": 4, "abr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "ago": 8,
          "sep": 9, "sept": 9, "set": 9, "oct": 10, "nov": 11, "dec": 12, "dic": 12}
ACCENTS = str.maketrans("áéíóúüñÁÉÍÓÚÜÑ", "aeiouunAEIOUUN")
PROVENANCE = ["source_file", "source_sheet", "source_row", "source_sha256", "extracted_at"]


def norm(s):
    """Lower case, accents stripped, punctuation turned to spaces."""
    return re.sub(r"[^a-z0-9 ]", " ", str(s or "").translate(ACCENTS).lower()).strip()


def period_from_tab(name):
    """'Mar 2026' -> 2026-03, 'Marzo' -> M03, '2026-03' -> 2026-03, '03' -> M03, else None."""
    n = norm(name)
    m = re.search(r"(20\d\d)\s*[- _/]?\s*(0?[1-9]|1[0-2])\b", n) or re.search(r"\b(0?[1-9]|1[0-2])\s*[- _/]?\s*(20\d\d)\b", n)
    if m:
        y, mo = (m.group(1), m.group(2)) if m.group(1).startswith("20") else (m.group(2), m.group(1))
        return f"{y}-{int(mo):02d}"
    for tok in n.split():
        key = tok[:4] if tok.startswith("sept") else tok[:3]
        if key in MONTHS and len(tok) <= 10:
            y = re.search(r"\b(20\d\d)\b", n)
            return f"{y.group(1)}-{MONTHS[key]:02d}" if y else f"M{MONTHS[key]:02d}"
    m = re.fullmatch(r"(0?[1-9]|1[0-2])", n)
    return f"M{int(m.group(1)):02d}" if m else None


def find_header(rows, columns, required, scan=30):
    """The row (within the first `scan`) whose cells match every required column, and
    where each canonical column sits. When several rows qualify, the one matching most
    columns wins."""
    syn = {c: {norm(s) for s in names} for c, names in columns.items()}
    best = (None, {})
    for i, row in enumerate(rows[:scan]):
        cells = [norm(v) for v in row]
        found = {}
        for canon, names in syn.items():
            for j, cell in enumerate(cells):
                if cell and (cell in names or any(cell.startswith(s) for s in names if len(s) > 3)) and j not in found.values():
                    found[canon] = j
                    break
        if all(c in found for c in required) and len(found) > len(best[1]):
            best = (i, found)
    return best


def sheets(path, only):
    """(tab name, rows) for each tab of a workbook, or the one table of a CSV file."""
    if path.suffix.lower() == ".csv":
        with open_csv_text(path) as fh:
            yield path.stem, [list(r) for r in csv.reader(fh)]
        return
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        if not only or name == only:
            yield name, [list(r) for r in wb[name].iter_rows(values_only=True)]


def spec_for(extract, extra):
    """The built-in spec for one extract type with any owner synonyms added."""
    spec = {"required": EXTRACTS[extract]["required"], "columns": {c: list(v) for c, v in EXTRACTS[extract]["columns"].items()}}
    for col, names in (extra.get(extract) or {}).items():
        spec["columns"].setdefault(col, []).extend(names)
    return spec


def normalize(path, extract, only_sheet=None, extra=None):
    spec = spec_for(extract, extra or {})
    digest = sha256(path)
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    out, skipped, used = [], [], []
    for sheet, rows in sheets(path, only_sheet):
        hdr, cols = find_header(rows, spec["columns"], spec["required"])
        if hdr is None:
            skipped.append(sheet)
            continue
        period = period_from_tab(sheet)
        footer = 0
        for i, row in enumerate(rows[hdr + 1:], start=hdr + 2):
            if not any(v not in (None, "") for v in row):
                continue
            rec = {c: (row[j] if j < len(row) else None) for c, j in cols.items()}
            if any(rec.get(c) in (None, "") for c in spec["required"]):
                footer += 1  # a subtotal or footer: a required key is blank
                continue
            if period and rec.get("period") in (None, ""):
                rec["period"] = period
            rec.update({"source_file": path.name, "source_sheet": sheet, "source_row": i, "source_sha256": digest[:16], "extracted_at": stamp})
            out.append(rec)
        used.append({"sheet": sheet, "header_row": hdr + 1, "columns": cols, "period_from_tab": period, "rows_skipped": footer})
    if not used:
        raise ValueError(f"DATA REQUEST: no tab in {path.name} has the headers for {extract!r}: {', '.join(spec['required'])}. Tabs seen: {', '.join(skipped)}")
    return {"extract": extract, "file": str(path), "sha256": digest, "tabs_used": used, "tabs_skipped": skipped, "rows": out}


def load_synonyms(path):
    path = path or settings(SKILL).get("extract_synonyms")
    if not path:
        return {}
    try:
        return json.loads(Path(path).expanduser().read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        fail(f"Could not read the synonyms file {path}: {exc}", 2)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Normalise one SAP extract file into a long table.")
    ap.add_argument("extract", nargs="?", help="Extract type (see --types)")
    ap.add_argument("path", nargs="?", help="The .xlsx or .csv file")
    ap.add_argument("--sheet", default="", help="Only this tab")
    ap.add_argument("--out", help="Write the long table as CSV (default: JSON to stdout)")
    ap.add_argument("--synonyms", help="JSON of extra header synonyms (default: the extract_synonyms setting)")
    ap.add_argument("--types", action="store_true", help="List extract types and their canonical columns")
    a = ap.parse_args(argv)

    if a.types or not a.extract or not a.path:
        for k, spec in EXTRACTS.items():
            print(f"{k:16} {', '.join(spec['columns'])}   (required: {', '.join(spec['required'])})")
        if not a.types:
            ap.error("EXTRACT and PATH are required")
        return 0
    if a.extract not in EXTRACTS:
        fail(f"Unknown extract type {a.extract!r}; one of {', '.join(EXTRACTS)}", 2)
    path = Path(a.path).expanduser()
    if not path.is_file():
        fail(f"Not a file: {path}", 2)
    try:
        r = normalize(path, a.extract, a.sheet or None, load_synonyms(a.synonyms))
    except ValueError as exc:
        fail(str(exc))
    if not a.out:
        print(json.dumps(r, indent=1, default=str))
        return 0
    cols = list(EXTRACTS[a.extract]["columns"])
    cols += [] if "period" in cols else ["period"]
    with Path(a.out).expanduser().open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols + PROVENANCE, extrasaction="ignore")
        w.writeheader()
        w.writerows(r["rows"])
    print(f"{len(r['rows'])} rows from {len(r['tabs_used'])} tabs -> {a.out}  (skipped: {', '.join(r['tabs_skipped']) or 'none'})  sha256 {r['sha256'][:12]}")
    for t in r["tabs_used"]:
        print(f"  {t['sheet']:24} header row {t['header_row']:3}  period {t['period_from_tab'] or '-':8}  {', '.join(t['columns'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
