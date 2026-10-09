# /// script
# dependencies = ["python-pptx"]
# ///
"""Build a board deck as a new .pptx from a deck spec.

The spec is JSON the writer keeps beside the package (work/deck-spec.json):

    {"slides": [
      {"layout": "title", "title": "Acme Components", "subtitle": "September 2026 board package"},
      {"title": "September EBITDA of $(74)k was $(56)k below forecast",
       "bullets": ["Revenue $951k, $(139)k below forecast", ["Development fees $(172)k"]],
       "table": {"columns": ["$k", "ACT", "FCST"], "rows": [["Revenue", "951", "1,090"]]},
       "notes": "Speaker notes for the commentary."}]}

A slide has a title and any of bullets (a nested list is the level below), table (columns and
rows of strings, printed exactly as given) and notes; layout "title" makes a title slide with
a subtitle. The renderer formats nothing, so the figure ledger and the deck cannot disagree.

--template builds on the company's own .pptx (its masters and layouts; its slides are left
out). An existing output is refused unless --supersede DIR is given, which first moves it to
DIR/superseded-<stamp>/ so every earlier draft is kept. Exit 0 written, 1 refused, 2 on a bad
spec or argument.

    python3 board_package_render.py work/deck-spec.json --out "Board Package 2026-09 DRAFT.pptx" --template template.pptx
"""

import argparse
import shutil
from pathlib import Path

import _common as bc


def load_spec(path: Path) -> dict:
    spec = bc.load_json(path, "deck spec")
    if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list) or not spec["slides"]:
        raise bc.Bad(f"{path.name}: a deck spec is an object with a non-empty 'slides' list")
    for n, slide in enumerate(spec["slides"], start=1):
        if not isinstance(slide, dict):
            raise bc.Bad(f"slide {n} is not an object")
        if not bc.blank(slide.get("title")):
            raise bc.Bad(f"slide {n} has no title")
        table = slide.get("table")
        if table is not None:
            cols, rows = table.get("columns"), table.get("rows")
            if not isinstance(cols, list) or not isinstance(rows, list) or not cols:
                raise bc.Bad(f"slide {n}: a table has 'columns' and 'rows' lists")
            for r_no, row in enumerate(rows, start=1):
                if not isinstance(row, list) or len(row) != len(cols):
                    raise bc.Bad(f"slide {n}: table row {r_no} has {len(row) if isinstance(row, list) else 0} "
                                 f"cells, not {len(cols)}")
        if slide.get("bullets") is not None and not isinstance(slide["bullets"], list):
            raise bc.Bad(f"slide {n}: 'bullets' is a list")
    return spec


def layout(prs, names, fallback):
    """The layout with one of these names, else the fallback position."""
    for item in prs.slide_layouts:
        if item.name.strip().lower() in names:
            return item
    layouts = list(prs.slide_layouts)
    return layouts[min(fallback, len(layouts) - 1)]


def bullet_lines(bullets, level=0) -> list:
    out = []
    for item in bullets:
        out += bullet_lines(item, level + 1) if isinstance(item, list) else [(str(item), level)]
    return out


def build(spec: dict, template):
    from pptx import Presentation
    from pptx.util import Emu, Pt

    if template is not None:
        if not template.is_file():
            raise bc.Bad(f"no deck template at {template}")
        prs = Presentation(str(template))
        ids = prs.slides._sldIdLst  # python-pptx has no public way to drop the template's slides
        for sld in list(ids):
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    else:
        prs = Presentation()
    width, height = prs.slide_width, prs.slide_height
    margin = Emu(int(width * 0.05))
    for spec_slide in spec["slides"]:
        if spec_slide.get("layout") == "title":
            slide = prs.slides.add_slide(layout(prs, ["title slide", "title"], 0))
            if slide.shapes.title is not None:
                slide.shapes.title.text = str(spec_slide["title"])
            subtitle = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1]
            if subtitle and spec_slide.get("subtitle"):
                subtitle[0].text = str(spec_slide["subtitle"])
        else:
            slide = prs.slides.add_slide(layout(prs, ["title only"], 5))
            if slide.shapes.title is not None:
                slide.shapes.title.text = str(spec_slide["title"])
            for ph in list(slide.placeholders):          # drop unused body placeholders
                if ph.placeholder_format.idx != 0:
                    ph._element.getparent().remove(ph._element)
            top = Emu(int(height * 0.22))
            lines = bullet_lines(spec_slide.get("bullets") or [])
            if lines:
                box_height = Emu(int(height * (0.30 if spec_slide.get("table") else 0.68)))
                frame = slide.shapes.add_textbox(margin, top, width - 2 * margin, box_height).text_frame
                frame.word_wrap = True
                for i, (text, level) in enumerate(lines):
                    para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    para.text = ("\u2022 " if level == 0 else "\u2013 ") + text
                    para.level = min(level, 4)
                    for run in para.runs:
                        run.font.size = Pt(16 if level == 0 else 14)
                top = Emu(top + box_height)
            table = spec_slide.get("table")
            if table:
                rows = [table["columns"]] + table["rows"]
                row_h = Emu(int(height * 0.055))
                shape = slide.shapes.add_table(len(rows), len(table["columns"]), margin, top,
                                               width - 2 * margin, Emu(row_h * len(rows)))
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        cell = shape.table.cell(r, c)
                        cell.text = "" if value is None else str(value)
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                run.font.size = Pt(12)
        if bc.blank(spec_slide.get("notes")):
            slide.notes_slide.notes_text_frame.text = str(spec_slide["notes"])
    return prs


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a deck from SPEC.json as a new .pptx; never overwrites.")
    parser.add_argument("spec", nargs="?", default="", metavar="SPEC")
    parser.add_argument("--out", default="", help="The new .pptx (refused if it exists)")
    parser.add_argument("--template", default="", help="The company's .pptx template")
    parser.add_argument("--supersede", default="", help="Move an existing output to DIR/superseded-<stamp>/ first")
    args = parser.parse_args()

    if not bc.blank(args.spec):
        raise bc.Bad("SPEC is required: the deck spec JSON")
    if not bc.blank(args.out) or not args.out.strip().lower().endswith(".pptx"):
        raise bc.Bad("--out is required and must end .pptx")
    out = Path(args.out.strip()).expanduser()
    spec = load_spec(Path(args.spec).expanduser())
    prs = build(spec, Path(args.template).expanduser() if bc.blank(args.template) else None)
    if out.exists():
        if not bc.blank(args.supersede):
            raise bc.Refused(f"{out} exists; the renderer never overwrites (pass --supersede WORKDIR "
                             f"to keep the earlier draft in WORKDIR/superseded-<stamp>/)")
        aside = Path(args.supersede.strip()).expanduser() / f"superseded-{bc.stamp()}"
        aside.mkdir(parents=True, exist_ok=True)
        shutil.move(str(out), str(aside / out.name))
        print(f"MOVED the earlier draft to {aside}")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"WROTE {out} ({len(spec['slides'])} slides)")
    return 0


if __name__ == "__main__":
    bc.run_main(main)
