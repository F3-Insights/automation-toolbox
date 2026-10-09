# /// script
# dependencies = ["openpyxl"]
# ///
"""Read an Excel workbook (.xlsx or .xlsm) without ever saving it.

The workbook is opened read-only with the values Excel last calculated, so a
formula cell shows its result, not its formula. The first row of a sheet is
taken as the column headers.

Operations (first argument):
  analyze   sheet names, row and column counts, per-column fill and sample values
  read      the rows of one sheet, optionally only some columns or the first N rows
  extract   the rows where one column meets a condition (">10", "==Open",
            "contains:text"); with no condition, the rows where it is not blank
  convert   one sheet (or every sheet with --all-sheets) to CSV files
  process   one column's unique values, statistics, or value frequencies

Output is JSON on stdout. read and extract print at most --limit rows (default
200; --limit 0 prints every row); when rows were cut, the output says how many
were shown out of how many in `rows_shown` and `note`. --output writes the result
to a file as well, always with every row: .csv or .json for read and extract,
.json for analyze and process, and the CSV path (or base name, with
--all-sheets) for convert.

Example:
  python3 excel_handle.py extract ledger.xlsx --sheet Detail --column Amount --condition ">1000"
"""

import argparse
import csv
import json
import statistics
import sys
from collections import Counter
from pathlib import Path

CONDITION_OPERATORS = (">=", "<=", "==", "!=", ">", "<")
DEFAULT_LIMIT = 200


def fail(message, code=1):
    print(message, file=sys.stderr)
    sys.exit(code)


def open_workbook(path):
    from openpyxl import load_workbook

    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        fail(f"not an .xlsx or .xlsm workbook: {path}")
    if not Path(path).exists():
        fail(f"file not found: {path}")
    return load_workbook(path, read_only=True, data_only=True)


def sheet_rows(workbook, sheet):
    """The sheet as a list of dicts keyed by the header row, blank rows skipped."""
    if sheet is None:
        sheet = workbook.sheetnames[0]
    if sheet not in workbook.sheetnames:
        fail(f"sheet '{sheet}' not found; sheets are {workbook.sheetnames}")
    rows = workbook[sheet].iter_rows(values_only=True)
    header = next(rows, ())
    columns = []
    for i, h in enumerate(header):
        name = str(h) if h is not None else f"Unnamed: {i}"
        # a repeated header gets a suffix (Amount, Amount.1) so no column's values are lost
        base, n = name, 0
        while name in columns:
            n += 1
            name = f"{base}.{n}"
        columns.append(name)
    records = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        records.append({c: (row[i] if i < len(row) else None) for i, c in enumerate(columns)})
    return sheet, columns, records


def need_column(columns, column):
    if not column:
        fail("--column is required for this operation")
    if column not in columns:
        fail(f"column '{column}' not found; columns are {columns}")


def matches(value, condition):
    """True when one cell meets a condition such as '>10', '==Open' or 'contains:net'."""
    if condition.startswith("contains:"):
        return value is not None and condition[9:].lower() in str(value).lower()
    for op in CONDITION_OPERATORS:
        if condition.startswith(op):
            target = condition[len(op):].strip()
            try:
                number = float(target)
            except ValueError:
                if op not in ("==", "!="):
                    fail(f"operator {op} needs a number, got '{target}'")
                equal = value is not None and str(value) == target
                return equal if op == "==" else not equal
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                return op == "!="
            return {">=": value >= number, "<=": value <= number, "==": value == number,
                    "!=": value != number, ">": value > number, "<": value < number}[op]
    fail(f"condition must start with one of {CONDITION_OPERATORS} or 'contains:': {condition}")


def write_rows(path, columns, records):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)


def save(path, result, columns=None, records=None):
    if not path:
        return
    if path.lower().endswith(".csv") and records is not None:
        write_rows(path, columns, records)
    elif path.lower().endswith(".json"):
        Path(path).write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    else:
        fail("--output must end in .csv or .json")


def analyze(workbook, args):
    sheet, columns, records = sheet_rows(workbook, args.sheet)
    info = {}
    for column in columns:
        values = [r[column] for r in records if r[column] is not None]
        info[column] = {
            "non_null_count": len(values),
            "null_count": len(records) - len(values),
            "unique_count": len({str(v) for v in values}),
            "sample_values": values[:5],
        }
    return {"file": args.input_file, "sheets": workbook.sheetnames, "active_sheet": sheet,
            "rows": len(records), "columns": len(columns), "column_info": info,
            "preview": records[:10]}, None, None


def read(workbook, args):
    sheet, columns, records = sheet_rows(workbook, args.sheet)
    if args.columns:
        wanted = [c.strip() for c in args.columns.split(",")]
        missing = [c for c in wanted if c not in columns]
        if missing:
            fail(f"columns not found: {missing}")
        columns = wanted
        records = [{c: r[c] for c in wanted} for r in records]
    if args.rows:
        records = records[:args.rows]
    return {"active_sheet": sheet, "rows": len(records), "columns": columns,
            "data": records}, columns, records


def extract(workbook, args):
    sheet, columns, records = sheet_rows(workbook, args.sheet)
    need_column(columns, args.column)
    if args.condition:
        kept = [r for r in records if matches(r[args.column], args.condition)]
    else:
        kept = [r for r in records if r[args.column] is not None]
    unique = list(dict.fromkeys(r[args.column] for r in kept))
    return {"active_sheet": sheet, "column": args.column, "condition": args.condition,
            "original_rows": len(records), "rows": len(kept), "unique_values": unique,
            "data": kept}, columns, kept


def convert(workbook, args):
    sheets = workbook.sheetnames if args.all_sheets else [args.sheet or workbook.sheetnames[0]]
    base = Path(args.output or args.input_file)
    written = []
    for sheet in sheets:
        _, columns, records = sheet_rows(workbook, sheet)
        if args.all_sheets:
            target = base.with_name(f"{base.stem}_{sheet}.csv")
        else:
            target = Path(args.output) if args.output else base.with_suffix(".csv")
        write_rows(target, columns, records)
        written.append(str(target))
    return {"sheets": workbook.sheetnames, "converted_files": written}, None, None


def process(workbook, args):
    sheet, columns, records = sheet_rows(workbook, args.sheet)
    need_column(columns, args.column)
    values = [r[args.column] for r in records if r[args.column] is not None]
    result = {"active_sheet": sheet, "column": args.column, "operation": args.col_operation,
              "non_null_count": len(values)}
    if args.col_operation == "unique":
        result["unique_values"] = list(dict.fromkeys(values))
        result["unique_count"] = len(result["unique_values"])
    elif args.col_operation == "stats":
        numbers = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if not numbers or len(numbers) != len(values):
            fail(f"column '{args.column}' is not numeric")
        result.update(mean=statistics.fmean(numbers), min=min(numbers), max=max(numbers),
                      median=statistics.median(numbers),
                      std=statistics.stdev(numbers) if len(numbers) > 1 else 0.0)
    else:
        counts = Counter(str(v) for v in values)
        result["frequency"] = dict(counts.most_common())
        result["most_common"] = dict(counts.most_common(10))
    return result, None, None


def limit_rows(result, limit):
    """Cut the printed rows to `limit` (0 means all) and say so; `rows` stays the full count."""
    data = result.get("data")
    if not limit or not isinstance(data, list) or len(data) <= limit:
        return result
    return dict(result, data=data[:limit], rows_shown=limit,
                note=(f"showing the first {limit} of {len(data)} rows; pass --limit 0 to print every "
                      "row, or --output to write them all to a file"))


OPERATIONS = {"analyze": analyze, "read": read, "extract": extract, "convert": convert,
              "process": process}


def main(argv=None):
    parser = argparse.ArgumentParser(description="Read an Excel workbook without saving it.")
    parser.add_argument("operation", choices=OPERATIONS)
    parser.add_argument("input_file")
    parser.add_argument("--sheet", "-s", help="sheet name (default: the first sheet)")
    parser.add_argument("--output", "-o", help="also write the result to this file")
    parser.add_argument("--rows", "-r", type=int, help="read: keep only the first N rows")
    parser.add_argument("--columns", "-c", help="read: comma-separated columns to keep")
    parser.add_argument("--column", help="extract and process: the column to work on")
    parser.add_argument("--condition", help='extract: ">10", "==value", "contains:text"')
    parser.add_argument("--all-sheets", action="store_true", help="convert: every sheet")
    parser.add_argument("--col-operation", choices=["unique", "stats", "frequency"],
                        default="unique", help="process: what to compute (default unique)")
    parser.add_argument("--limit", type=int, default=DEFAULT_LIMIT,
                        help=f"read and extract: print at most N rows (default {DEFAULT_LIMIT}; 0 prints all)")
    args = parser.parse_args(argv)
    if args.limit < 0:
        parser.error("--limit must be 0 or more")

    workbook = open_workbook(args.input_file)
    result, columns, records = OPERATIONS[args.operation](workbook, args)
    if args.operation != "convert":
        save(args.output, result, columns, records)
    print(json.dumps(limit_rows(result, args.limit), indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
