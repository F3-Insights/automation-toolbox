"""Reusable python-pptx layout primitives for the seminar-builder skill.

Per-seminar build scripts import these helpers to construct branded slides.
The style tokens below mirror `references/deck-styling.md` - keep both in sync
when rebranding.

These helpers were extracted from a working per-seminar build script.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Style tokens - F3 Insights brand palette
#   Inter font (Google) - falls back to Calibri/system-ui if not installed
#   Edit references/deck-styling.md and these constants together when rebranding.
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1E, 0x3A, 0x6E)          # Primary text, hero bands
BRONZE = RGBColor(0xA6, 0x7C, 0x52)        # Accent: subtitles, top chrome, callouts
BRONZE_LIGHT = RGBColor(0xB8, 0x8B, 0x5E)  # Secondary accent (ring/hover)
GRAY = RGBColor(0x4B, 0x55, 0x63)          # Muted body text
LIGHT_GRAY = RGBColor(0xF9, 0xFA, 0xFB)    # Card / panel backgrounds
BORDER_GRAY = RGBColor(0xE5, 0xE7, 0xEB)   # Borders, dividers
CODE_BG = RGBColor(0x1F, 0x29, 0x37)       # Code block background
CODE_FG = RGBColor(0xE6, 0xED, 0xF3)       # Code text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_RED = RGBColor(0xB4, 0x1E, 0x1E)    # Warnings, "bad" column
ACCENT_GREEN = RGBColor(0x1A, 0x7F, 0x37)  # Success, "good" column

# Alias kept so call sites that reference ORANGE still work
ORANGE = BRONZE

TITLE_FONT = "Inter"
BODY_FONT = "Inter"
CODE_FONT = "Consolas"

SLIDE_W = Inches(13.333)  # 16:9 widescreen
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)


# ---------------------------------------------------------------------------
# Slide construction
# ---------------------------------------------------------------------------
def new_presentation() -> Presentation:
    """Create a 16:9 widescreen presentation sized to SLIDE_W x SLIDE_H."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_rect(slide, left, top, width, height, fill_rgb, line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if not line:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = fill_rgb
    shape.shadow.inherit = False
    return shape


def add_text(
    slide,
    text: str,
    left,
    top,
    width,
    height,
    *,
    font=BODY_FONT,
    size=18,
    bold=False,
    italic=False,
    color=NAVY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor

    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_bullets(slide, items, left, top, width, height, *, size=18, color=NAVY, bullet="•"):
    """Items are strings or (head, sub) tuples. Sub-text indents under head."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)

        if isinstance(item, tuple):
            head, sub = item
        else:
            head, sub = item, None

        r = p.add_run()
        r.text = f"{bullet}  {head}"
        r.font.name = BODY_FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = True

        if sub:
            sp = tf.add_paragraph()
            sp.alignment = PP_ALIGN.LEFT
            sp.space_after = Pt(10)
            sp.level = 1
            sr = sp.add_run()
            sr.text = sub
            sr.font.name = BODY_FONT
            sr.font.size = Pt(size - 4)
            sr.font.color.rgb = GRAY
            sr.font.bold = False
    return box


def add_code(slide, code: str, left, top, width, height, *, size=12):
    """Dark code block: monospace text on CODE_BG background."""
    add_rect(slide, left, top, width, height, CODE_BG)
    box = slide.shapes.add_textbox(
        left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), height - Inches(0.2)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line if line else " "
        r.font.name = CODE_FONT
        r.font.size = Pt(size)
        r.font.color.rgb = CODE_FG


def add_chrome(slide, slide_num: int, total: int, section: str = ""):
    """Top accent bar + bottom footer with slide number. Call on every non-title slide."""
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.15), BRONZE)
    add_text(
        slide,
        f"{section}",
        MARGIN,
        SLIDE_H - Inches(0.4),
        Inches(8),
        Inches(0.3),
        size=10,
        color=GRAY,
        italic=True,
    )
    add_text(
        slide,
        f"{slide_num} / {total}",
        SLIDE_W - Inches(1.2),
        SLIDE_H - Inches(0.4),
        Inches(0.8),
        Inches(0.3),
        size=10,
        color=GRAY,
        align=PP_ALIGN.RIGHT,
    )


def slide_title(slide, title: str, subtitle: str | None = None):
    """Standard title + optional bronze italic subtitle at top of slide."""
    add_text(
        slide,
        title,
        MARGIN,
        Inches(0.45),
        SLIDE_W - 2 * MARGIN,
        Inches(0.8),
        font=TITLE_FONT,
        size=32,
        bold=True,
        color=NAVY,
    )
    if subtitle:
        add_text(
            slide,
            subtitle,
            MARGIN,
            Inches(1.25),
            SLIDE_W - 2 * MARGIN,
            Inches(0.5),
            size=18,
            color=BRONZE,
            italic=True,
        )


def set_notes(slide, script_text: str) -> None:
    """Embed a speaker script in the PowerPoint notes pane for this slide.

    Call once per slide with the 60–120s script from the Step-3 outline.
    The notes pane shows under the slide in Presenter View.
    """
    slide.notes_slide.notes_text_frame.text = script_text
