# /// script
# dependencies = ["openpyxl"]
# ///
"""Audit a per-finished-good cost model workbook without editing it.

Reads the workbook through the same layout file as fg_pl_extract.py (--layout or the
`cost_model_layout` setting in [product-costing-workstream]); tab names can be overridden
on the command line. The report carries the file's SHA-256 so each finding ties to the
exact file. Checks, each a finding with a severity:

  cached_errors      (error) any cell whose saved value is #REF!, #DIV/0! and the like
  duplicate_keys     (error) the same product twice for one month on the month tab
  margin_arithmetic  (error) cm = sales - total cost within $1; cm_pct = cm / sales; a
                     product with revenue and no cost shows a blank margin, not 0% or 100%
  ytd_equals_sum     (error) YTD block N equals months 1..N for sales and total cost
  std_var_ties_tb    (error) for the audited month, standard plus variance per cost
                     component equals the TB for that category within max($1, 0.05%);
                     --tb-map JSON {"dm": ["Materials"], ...} names the TB categories,
                     otherwise the component's own name (dm, dl, ...) is matched
  wrong_month_refs   (error) formulas on the allocation tab that point at a TB month
                     column or a month block other than the audited month
  factory_pools_sum  (error) with --pool-tabs and --consolidated-tab, the per-factory TB
                     tabs add up to the consolidated TB per category and month

Prints a text report (or JSON with --format json). Exit 1 on any error finding, 2 when
the layout is missing.

Example:
  python3 costmodel_audit.py model.xlsx --year 2026 --month 8 --layout layout.json \\
      --pool-tabs "TB plant A,TB plant B" --consolidated-tab "TB total"
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from _common import ERRS, fail, find_tab, load_layout, parse_product_tab, parse_tb, sha256

# 'Sheet Name'!$F$12  or  Sheet!F12
REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\.]+))?!?\$?([A-Z]{1,3})\$?(\d+)")


def col_index(letters):
    """Column letters to a 0-based index: A -> 0, AA -> 26."""
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def cost_components(layout):
    """The cost components the block carries as std_<x> and var_<x> pairs."""
    fields = layout["block_fields"]
    return [f[4:] for f in fields if f.startswith("std_") and f"var_{f[4:]}" in fields]


def check_cached_errors(wb):
    out = []
    for ws in wb.worksheets:
        n, first = 0, None
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRS:
                    n += 1
                    first = first or f"{c.coordinate}={c.value}"
        if n:
            out.append({"check": "cached_errors", "severity": "error", "tab": ws.title, "detail": f"{n} error cells, first {first}"})
    return out


def check_duplicates(fg):
    seen = defaultdict(int)
    for r in fg:
        if r["basis"] == "month":
            seen[(r["product_fg"], r["month_no"])] += 1
    dups = {}
    for (product, _month), n in seen.items():
        if n > 1:
            dups[product] = max(dups.get(product, 0), n)
    return [{"check": "duplicate_keys", "severity": "error", "product": k, "detail": f"{v} rows for the same product and month"} for k, v in sorted(dups.items())]


def check_margins(fg):
    out = []
    for r in fg:
        rev, cost, cm = r.get("sales_amount") or 0.0, r.get("total_cost") or 0.0, r.get("cm")
        where = f"{r['basis']} {r['month_no']}"
        if r.get("cost_pending"):
            # the parser blanks the percentage itself, so judge the value saved in the sheet
            if r.get("_raw_cm_pct") in (0.0, 1.0):
                out.append({"check": "margin_arithmetic", "severity": "error", "product": r["product_fg"], "detail": f"{where}: revenue with no cost shows {r['_raw_cm_pct']:.0%} margin; should be blank (cost pending)"})
            continue
        if cm is not None and abs((rev - cost) - cm) > 1.0:
            out.append({"check": "margin_arithmetic", "severity": "error", "product": r["product_fg"], "detail": f"{where}: cm {cm:,.2f} != sales {rev:,.2f} - cost {cost:,.2f}"})
        pct = r.get("cm_pct")
        if pct is not None and rev and cm is not None and abs(pct - cm / rev) > 0.0005:
            out.append({"check": "margin_arithmetic", "severity": "error", "product": r["product_fg"], "detail": f"{where}: cm_pct {pct:.4f} != cm/sales {cm / rev:.4f}"})
    return out


def check_ytd(fg):
    months, ytd = defaultdict(dict), defaultdict(dict)
    for r in fg:
        (months if r["basis"] == "month" else ytd)[r["product_fg"]][r["month_no"]] = r
    out = []
    for product, blocks in ytd.items():
        for n, yrow in blocks.items():
            for f in ("sales_amount", "total_cost"):
                total = sum((months[product].get(m, {}).get(f) or 0.0) for m in range(1, n + 1))
                y = yrow.get(f) or 0.0
                if abs(total - y) > 1.0:
                    out.append({"check": "ytd_equals_sum", "severity": "error", "product": product, "detail": f"YTD {n} {f} {y:,.2f} != sum of months {total:,.2f}"})
    return out


def check_std_var_tb(fg, tb, month, tb_map, components):
    out = []
    for cat in components:
        model = sum((r.get(f"std_{cat}") or 0.0) + (r.get(f"var_{cat}") or 0.0) for r in fg if r["basis"] == "month" and r["month_no"] == month)
        names = {n.lower() for n in tb_map.get(cat, [cat])}
        book = sum(r["amount"] for r in tb if r["month_no"] == month and ((r.get("category") or "").lower() in names or (r.get("subcategory") or "").lower() in names))
        if not model and not book:
            continue
        tol = max(1.0, 0.0005 * abs(book))
        if abs(model - book) > tol:
            out.append({"check": "std_var_ties_tb", "severity": "error", "category": cat, "detail": f"month {month}: model std+var {model:,.2f} vs TB {book:,.2f} (diff {model - book:,.2f}, tol {tol:,.2f})"})
    return out


def check_month_refs(wb_formulas, month, layout):
    alloc = layout.get("alloc_tab")
    if not alloc or alloc not in wb_formulas.sheetnames:
        return [{"check": "wrong_month_refs", "severity": "info", "tab": alloc or "-", "detail": "allocation tab not found; check skipped"}]
    tb_prefix, product_tab = layout["tb_tab"].lower(), layout["month_tab"]
    start, width, tb_first = layout["block_start"], layout["block_width"], layout["tb_first_amount_col"]
    bad = defaultdict(list)
    for row in wb_formulas[alloc].iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            for quoted, bare, col, _row in REF_RE.findall(c.value):
                sheet = quoted or bare
                if not sheet:
                    continue
                idx = col_index(col)
                if sheet.lower().startswith(tb_prefix):
                    ref = idx - tb_first + 1
                    if 1 <= ref <= layout["periods"] and ref != month:
                        bad[f"{sheet} month {ref}"].append(c.coordinate)
                elif sheet == product_tab and idx >= start:
                    ref = (idx - start) // width + 1
                    if ref != month:
                        bad[f"{sheet} block {ref}"].append(c.coordinate)
    return [{"check": "wrong_month_refs", "severity": "error", "tab": alloc, "detail": f"{len(cells)} formulas reference {target} instead of month {month} (first {cells[0]})"} for target, cells in sorted(bad.items())]


def check_pools(wb, pool_tabs, consolidated, year, layout):
    missing = [t for t in pool_tabs + [consolidated] if t not in wb.sheetnames]
    if missing:
        return [{"check": "factory_pools_sum", "severity": "error", "detail": f"tabs not found: {missing}"}]

    def totals(tab):
        agg = defaultdict(float)
        for r in parse_tb(wb[tab], year, layout):
            agg[((r.get("category") or "").lower(), r["month_no"])] += r["amount"]
        return agg

    cons, pools = totals(consolidated), defaultdict(float)
    for t in pool_tabs:
        for k, v in totals(t).items():
            pools[k] += v
    out = []
    for key in sorted(set(cons) | set(pools)):
        c, p = cons.get(key, 0.0), pools.get(key, 0.0)
        if abs(c - p) > max(1.0, 0.0005 * abs(c)):
            out.append({"check": "factory_pools_sum", "severity": "error", "category": key[0], "detail": f"month {key[1]}: pools {p:,.2f} vs consolidated {c:,.2f} (diff {p - c:,.2f})"})
    return out


def raw_margin_pcts(ws, layout):
    """The cm_pct value saved in the sheet, per (product, month), before the parser blanks it."""
    pid = layout["info_columns"]["product_fg"]
    col = layout["block_start"] + layout["block_fields"].index("cm_pct")
    out = {}
    for r in ws.iter_rows(min_row=layout["first_row"], values_only=True):
        if r and len(r) > pid and r[pid]:
            for b in range(layout["periods"]):
                i = col + layout["block_width"] * b
                if i < len(r):
                    out[(str(r[pid]).strip(), b + 1)] = r[i]
    return out


def audit(path, year, month, layout, pool_tabs=(), consolidated="", tb_map=None):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    tb_name = find_tab(wb.sheetnames, layout["tb_tab"])
    if layout["month_tab"] not in wb.sheetnames or tb_name is None:
        raise ValueError(f"Need tabs {layout['month_tab']!r} and a TB tab starting {layout['tb_tab']!r}; sheets: {wb.sheetnames}")
    fg = parse_product_tab(wb[layout["month_tab"]], "month", year, layout)
    raw = raw_margin_pcts(wb[layout["month_tab"]], layout)
    for r in fg:
        r["_raw_cm_pct"] = raw.get((r["product_fg"], r["month_no"]))
    if layout.get("ytd_tab") in wb.sheetnames:
        fg += parse_product_tab(wb[layout["ytd_tab"]], "ytd", year, layout)
    tb = parse_tb(wb[tb_name], year, layout)

    findings = check_cached_errors(wb) + check_duplicates(fg) + check_margins(fg) + check_ytd(fg)
    findings += check_std_var_tb(fg, tb, month, tb_map or {}, cost_components(layout))
    findings += check_month_refs(openpyxl.load_workbook(path, data_only=False), month, layout)
    if pool_tabs and consolidated:
        findings += check_pools(wb, list(pool_tabs), consolidated, year, layout)
    errors = sum(1 for f in findings if f["severity"] == "error")
    return {"workbook": str(path), "sha256": sha256(path), "year": year, "month": month, "tabs": wb.sheetnames,
            "products": len({r["product_fg"] for r in fg}), "findings": findings, "errors": errors, "passed": errors == 0}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Audit a cost model workbook without editing it. Exit 1 on any error finding.")
    ap.add_argument("workbook")
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--month", type=int, required=True, help="The month this version reports (1-12)")
    ap.add_argument("--layout", help="Layout JSON (default: the cost_model_layout setting)")
    ap.add_argument("--month-tab", help="Override the layout's month tab")
    ap.add_argument("--ytd-tab", help="Override the layout's year-to-date tab")
    ap.add_argument("--tb-tab", help="Override the layout's trial-balance tab (prefix match)")
    ap.add_argument("--alloc-tab", help="Override the layout's allocation tab")
    ap.add_argument("--pool-tabs", default="", help="Comma-separated per-factory TB tabs")
    ap.add_argument("--consolidated-tab", default="", help="Consolidated TB tab the pools must sum to")
    ap.add_argument("--tb-map", default="", help="JSON file mapping cost components (dm, dl, ...) to TB category names")
    ap.add_argument("--format", choices=["text", "json"], default="text")
    a = ap.parse_args(argv)

    path = Path(a.workbook).expanduser()
    if not path.is_file():
        fail(f"Not a file: {path}")
    layout = load_layout(a.layout, {"month_tab": a.month_tab, "ytd_tab": a.ytd_tab, "tb_tab": a.tb_tab, "alloc_tab": a.alloc_tab})
    try:
        tb_map = json.loads(Path(a.tb_map).expanduser().read_text(encoding="utf-8")) if a.tb_map else {}
        pools = [t.strip() for t in a.pool_tabs.split(",") if t.strip()]
        r = audit(path, a.year, a.month, layout, pools, a.consolidated_tab, tb_map)
    except (OSError, ValueError) as exc:
        fail(str(exc))
    if a.format == "json":
        print(json.dumps(r, indent=1, default=str))
    else:
        print(f"{path.name}  sha256 {r['sha256'][:12]}  {r['products']} products  month {r['month']}")
        for f in r["findings"]:
            where = f.get("product") or f.get("tab") or f.get("category") or ""
            print(f"  {f['severity'].upper():7} {f['check']:20} {where:24} {f['detail']}")
        print(f"{len(r['findings'])} findings, {r['errors']} errors: " + ("PASS" if r["passed"] else "FAIL"))
    return 0 if r["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
