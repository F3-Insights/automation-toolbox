#!/usr/bin/env python3
"""Render each slide of a generated .pptx to a PNG preview.

This is part of the seminar-builder skill's Step 4 visual-verification loop.
After building a deck, render previews and visually inspect each slide for
cramped text, overlaps, off-balance layouts, or unpolished visuals before
declaring the deck done.

Reads the .pptx directly (single source of truth = build_deck.py's output),
walks shapes, and draws an approximation with Pillow. Not pixel-perfect with
PowerPoint - the goal is layout inspection (overlaps, alignment, density).

Usage:
    python render_previews.py <path-to-pptx>
    python render_previews.py             # auto-detects single .pptx in cwd

Output: slide-previews/slide-NN.png in the directory of the .pptx file.
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DPI = 120  # 13.33"x7.5" → 1600x900 px

FONTS = {
    (False, False): "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    (True, False): "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    (False, True): "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
    (True, True): "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf",
}
MONO_FONTS = {
    (False, False): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    (True, False): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf",
    (False, True): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Oblique.ttf",
    (True, True): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-BoldOblique.ttf",
}


def emu_to_px(emu: int) -> int:
    return int(round(emu / 914400 * DPI))


def font_for(name, size_pt, bold, italic):
    is_mono = name and ("Mono" in name or "Consolas" in name or "Courier" in name)
    table = MONO_FONTS if is_mono else FONTS
    path = table.get((bold, italic)) or table[(False, False)]
    pixels = int(round(size_pt * DPI / 72))
    try:
        return ImageFont.truetype(path, pixels)
    except OSError:
        return ImageFont.load_default()


def rgb_tuple(color_obj):
    try:
        rgb = color_obj.rgb
        if rgb is None:
            return None
        return (rgb[0], rgb[1], rgb[2])
    except (AttributeError, KeyError, TypeError, ValueError):
        return None


def shape_fill(shape):
    try:
        fill = shape.fill
        if fill.type == 1:  # MSO_FILL.SOLID
            return rgb_tuple(fill.fore_color)
    except (AttributeError, KeyError, ValueError):
        pass
    return None


def wrap_text(draw, text, font, max_w):
    if not text:
        return [""]
    words = text.split(" ")
    lines = []
    cur = ""
    for w in words:
        trial = w if not cur else f"{cur} {w}"
        bbox = draw.textbbox((0, 0), trial, font=font)
        if bbox[2] - bbox[0] <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def render_text_frame(draw, shape):
    tf = shape.text_frame
    x = emu_to_px(shape.left)
    y = emu_to_px(shape.top)
    w = emu_to_px(shape.width)
    h = emu_to_px(shape.height)

    anchor = tf.vertical_anchor or MSO_ANCHOR.TOP

    blocks = []
    for p in tf.paragraphs:
        text = "".join(r.text or "" for r in p.runs)
        if not p.runs:
            continue
        r0 = p.runs[0]
        size_pt = (r0.font.size.pt if r0.font.size else 18)
        bold = bool(r0.font.bold)
        italic = bool(r0.font.italic)
        name = r0.font.name
        font = font_for(name, size_pt, bold, italic)
        rgb = rgb_tuple(r0.font.color) or (28, 32, 42)
        align = p.alignment or PP_ALIGN.LEFT
        indent_px = 30 * (p.level or 0)
        avail_w = max(10, w - indent_px)
        lines = wrap_text(draw, text, font, avail_w)
        leading = int(size_pt * DPI / 72 * 1.2)
        space_after_px = int((p.space_after.pt if p.space_after else 0) * DPI / 72)
        blocks.append((lines, font, rgb, align, leading, indent_px, space_after_px))

    total_h = sum(len(lines) * leading + space_after_px
                  for lines, _, _, _, leading, _, space_after_px in blocks)

    if anchor == MSO_ANCHOR.MIDDLE:
        cur_y = y + max(0, (h - total_h) // 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        cur_y = y + max(0, h - total_h)
    else:
        cur_y = y

    for lines, font, rgb, align, leading, indent_px, space_after_px in blocks:
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font)
            line_w = bbox[2] - bbox[0]
            if align == PP_ALIGN.CENTER:
                lx = x + (w - line_w) // 2
            elif align == PP_ALIGN.RIGHT:
                lx = x + w - line_w
            else:
                lx = x + indent_px
            draw.text((lx, cur_y), line, font=font, fill=rgb)
            cur_y += leading
        cur_y += space_after_px


def render_slide(slide, slide_w_px, slide_h_px):
    img = Image.new("RGB", (slide_w_px, slide_h_px), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    for shape in slide.shapes:
        fill_rgb = shape_fill(shape)
        if fill_rgb is not None and shape.shape_type is not None:
            x = emu_to_px(shape.left)
            y = emu_to_px(shape.top)
            w = emu_to_px(shape.width)
            h = emu_to_px(shape.height)
            draw.rectangle([x, y, x + w, y + h], fill=fill_rgb)

        if shape.has_text_frame and shape.text_frame.text.strip():
            render_text_frame(draw, shape)

    return img


def main():
    # Resolve target .pptx
    if len(sys.argv) > 1:
        pptx_path = Path(sys.argv[1]).resolve()
    else:
        # Auto-detect single .pptx in cwd
        candidates = list(Path.cwd().glob("*.pptx"))
        if len(candidates) == 1:
            pptx_path = candidates[0].resolve()
        elif len(candidates) == 0:
            print("error: no .pptx file in current directory; pass path explicitly")
            sys.exit(1)
        else:
            print(f"error: multiple .pptx files in cwd; pass one as argument:")
            for c in candidates:
                print(f"  {c.name}")
            sys.exit(1)

    out_dir = pptx_path.parent / "slide-previews"
    out_dir.mkdir(exist_ok=True)
    for f in out_dir.glob("slide-*.png"):
        f.unlink()

    prs = Presentation(str(pptx_path))
    sw = emu_to_px(prs.slide_width)
    sh = emu_to_px(prs.slide_height)
    print(f"Rendering {len(prs.slides)} slides at {sw}x{sh} px from {pptx_path.name}")

    for i, slide in enumerate(prs.slides, 1):
        img = render_slide(slide, sw, sh)
        out = out_dir / f"slide-{i:02d}.png"
        img.save(out, optimize=True)
        print(f"  ✓ {out.name}  ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
